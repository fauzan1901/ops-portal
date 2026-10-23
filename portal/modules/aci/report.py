"""Export the static port report to Excel (.xlsx) and CSV."""
import csv
import datetime
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

COLUMNS = ["Site", "Tenant", "Application Profile", "EPG", "VLAN", "Binding Type", "Pod", "Node(s)",
           "Interface / Policy Group", "Mode", "Deployment Immediacy", "Encapsulation", "Path TDN",
           "Source", "Portal Change", "Portal Setting", "Last Changed (APIC)"]
WIDTHS = [8, 22, 28, 28, 8, 12, 6, 14, 30, 11, 20, 14, 70, 24, 14, 32, 20]
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
BINDING_TYPES = ("Physical", "PC", "vPC", "Other")
SOURCES = ("Portal", "Outside portal", "Changed outside portal", "Removed outside portal")


def apic_time(value):
    """APIC timestamp (2026-09-25T08:48:19.123+07:00) in the local time of this server: 2026-09-25 08:48:19."""
    if not value:
        return ""
    try:
        return datetime.datetime.fromisoformat(str(value).replace("Z", "+00:00")).astimezone().strftime("%Y-%m-%d %H:%M:%S")
    except ValueError:
        return str(value)


def column_value(row, column):
    """Value of one column; the portal columns come from what the portal last set on that path (row["portal"])."""
    portal = row.get("portal") or {}
    if column == "Portal Change":
        return portal.get("change", "")
    if column == "Portal Setting":
        if not portal:
            return ""
        return "deleted" if portal.get("status") == "delete" else ", ".join(str(portal.get(k) or "") for k in ("encap", "mode", "immediacy"))
    if column == "Last Changed (APIC)":
        return apic_time(row.get(column))
    return row.get(column, "")


def to_xlsx(rows, info):
    """info: dict of metadata for the Summary sheet (site, APIC, filter, user, time)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Static Port Report"
    ws.sheet_view.showGridLines = False
    ws.append(COLUMNS)
    for row in rows:
        ws.append([column_value(row, c) for c in COLUMNS])
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for i, width in enumerate(WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    last = get_column_letter(len(COLUMNS))
    if rows:
        # The Excel table already has its own filter; do not stack it with the sheet autofilter.
        table = Table(displayName="ACIStaticPortReport", ref=f"A1:{last}{ws.max_row}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(table)
    else:
        ws.auto_filter.ref = f"A1:{last}1"

    summary = wb.create_sheet("Summary")
    summary.sheet_view.showGridLines = False
    summary.append(["Metric", "Value"])
    for key, value in info.items():
        summary.append([key, value])
    summary.append(["Total Bindings", len(rows)])
    summary.append(["Tenants", len({r["Tenant"] for r in rows})])
    summary.append(["Application Profiles", len({(r["Tenant"], r["Application Profile"]) for r in rows})])
    summary.append(["EPGs", len({(r["Tenant"], r["Application Profile"], r["EPG"]) for r in rows})])
    for kind in BINDING_TYPES:
        summary.append([f"{kind} Bindings", sum(r["Binding Type"] == kind for r in rows)])
    for source in SOURCES:
        summary.append([f"Source: {source}", sum(r.get("Source") == source for r in rows)])
    for cell in summary[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    summary.column_dimensions["A"].width = 24
    summary.column_dimensions["B"].width = 48

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def safe_cell(value):
    """Prevent formula injection when the CSV is opened in Excel (text starting with = + - @ tab CR)."""
    if value is None:
        return ""
    if isinstance(value, (int, float)):
        return value
    text = str(value)
    return "'" + text if text[:1] in ("=", "+", "-", "@", "\t", "\r") else text


def to_csv(rows):
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow(COLUMNS)
    writer.writerows([safe_cell(column_value(row, c)) for c in COLUMNS] for row in rows)
    return out.getvalue()
