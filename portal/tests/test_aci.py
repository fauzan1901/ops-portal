"""Tests for the ACI module of the portal (/aci/): validation (unit) and the full flow against three mock APICs
(end-to-end), with sign-in on the portal through a fake Active Directory (LDAPS), a fake Duo and a fake Duo SSO, module
access per user, and one APIC service account per site of sites.yml.

Run from the portal folder:

    .venv/bin/python -m unittest tests.test_aci -v
"""
import csv
import datetime
import io
import json
import os
import re
import shutil
import sqlite3
import ssl
import sys
import tempfile
import time
import unittest
import unittest.mock
import urllib.error
import urllib.parse
import urllib.request

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))
sys.path.insert(0, HERE)

# The portal reads its settings when it is imported: every file goes to a temporary folder of these tests, and only the
# ACI module is loaded. sites.yml and auth.yml are written by EndToEndTest once the mock servers run.
TMP = tempfile.mkdtemp(prefix="portal-test-aci-")
ENV = {"PORTAL_SITES": os.path.join(TMP, "sites.yml"), "PORTAL_DATA": TMP, "PORTAL_DB": os.path.join(TMP, "portal.db"),
       "PORTAL_AUTH": os.path.join(TMP, "auth.yml"), "PORTAL_CREDENTIALS_KEY": os.path.join(TMP, "credentials.key"),
       "PORTAL_MODULES": "aci"}
SAVED_ENV = {name: os.environ.get(name) for name in ENV}
os.environ.update(ENV)

import app as portal  # noqa: E402
import fake_duo  # noqa: E402
import fake_sso  # noqa: E402
import fake_ldap  # noqa: E402
import modules.aci as aci_module  # noqa: E402
from mock_apic import MockApic, make_cert  # noqa: E402
from modules.aci import aci, audit, ip_usage, product_templates, template_store  # noqa: E402


def tearDownModule():
    for name, value in SAVED_ENV.items():
        if value is None:
            os.environ.pop(name, None)
        else:
            os.environ[name] = value
    shutil.rmtree(TMP, ignore_errors=True)


REFRESH_TIMEOUT = 4  # seconds; the app refreshes the token after 2 seconds
SERVICE_USER = "svc_aci"
AD_USERS = {"admin": {"password": "secret", "display_name": "Ada Admin"},
            "operator2": {"password": "secret", "display_name": "Otto Operator"},
            "operator3": {"password": "secret", "display_name": "Oscar Operator"},   # a second operator (undo permissions)
            "operator4": {"password": "secret", "display_name": "Lina L1"},          # on the portal, but without ACI
            "outsider": {"password": "secret", "display_name": "Olga Outsider"}}


class ValidationTest(unittest.TestCase):
    def test_check_name(self):
        self.assertEqual(aci.check_name(" OpenStack-Tenant ", "Tenant"), "OpenStack-Tenant")
        for bad in ("", "a/b", "x.json?y", "a b", "x" * 65):
            with self.assertRaises(ValueError):
                aci.check_name(bad, "Tenant")

    def test_bindings_accept_aliases(self):
        bindings = aci.normalize_bindings(
            [{"pod": "1", "node": "106", "interface": "Eth1/17", "vlan": "vlan-25", "mode": "trunk", "immediacy": "lazy"}],
            [{"tdn": "topology/pod-1/protpaths-103-104/pathep-[PG-A]", "vlan": "11", "mode": "access"}])
        self.assertEqual(bindings[0], {"tdn": "topology/pod-1/paths-106/pathep-[eth1/17]", "encap": "vlan-25", "mode": "regular", "immediacy": "lazy"})
        self.assertEqual(bindings[1]["mode"], "untagged")

    def test_bindings_collect_all_errors(self):
        with self.assertRaises(ValueError) as ctx:
            aci.normalize_bindings(
                [{"node": "101", "interface": "eth1/1", "vlan": "5000"}, {"node": "x", "interface": "eth1/2", "vlan": "10"},
                 {"node": "101", "interface": "eth1/3", "vlan": "10"}, {"node": "101", "interface": "eth1/3", "vlan": "12"}],
                [{"tdn": "uni/bad", "vlan": "10"}])
        message = str(ctx.exception)
        self.assertIn("Physical/PC row 1: VLAN must be 1-4094", message)
        self.assertIn("Physical/PC row 2: Node must be a number", message)
        self.assertIn("vPC/TDN row 1: Invalid path TDN", message)
        self.assertIn("entered more than once", message)

    def test_path_info(self):
        self.assertEqual(aci.path_info("topology/pod-1/paths-101/pathep-[eth1/1]")["Binding Type"], "Physical")
        self.assertEqual(aci.path_info("topology/pod-1/paths-101/pathep-[PC-SERVER-01]")["Binding Type"], "PC")
        fex = aci.path_info("topology/pod-1/paths-101/extpaths-111/pathep-[eth1/5]")
        self.assertEqual((fex["Binding Type"], fex["Node(s)"]), ("Physical", "101/FEX-111"))
        self.assertEqual(aci.path_info("topology/pod-1/protpaths-103-104/pathep-[PG]")["Node(s)"], "103-104")

    def test_subnet(self):
        self.assertEqual(aci.check_subnet({"ip": "10.253.0.1/24"})["scope"], "public")
        self.assertEqual(aci.check_subnet({"ip": "10.253.0.1/24", "scope": "private,shared"})["scope"], "private,shared")
        for bad in ({"ip": "10.253.0.0/24"}, {"ip": "10.253.0.1"}, {"ip": "10.253.0.1/24", "scope": "public,private"}, {"ip": "abc/24"}):
            with self.assertRaises(ValueError):
                aci.check_subnet(bad)


class ProductTemplateTest(unittest.TestCase):
    def test_expand_interfaces(self):
        self.assertEqual(product_templates.expand_interfaces("eth1/1-3, eth1/5, Po-Solarwinds"), ["eth1/1", "eth1/2", "eth1/3", "eth1/5", "Po-Solarwinds"])
        self.assertEqual(product_templates.expand_interfaces("eth1/1/1-2"), ["eth1/1/1", "eth1/1/2"])
        for bad in ("eth1/5-3", "eth1/[1]", "a b"):
            with self.assertRaises(ValueError):
                product_templates.expand_interfaces(bad)

    def test_shipped_templates(self):
        # Number of default paths per site, from the analysis of the ACI Report of each site (physical ports + vPC).
        expected = {("DCB", "flex"): 14, ("DCB", "prime"): 25, ("DCB", "rhoso"): 7, ("DCA", "flex"): 11, ("DCA", "prime"): 46, ("DCA", "rhoso"): 0,
                    ("DCC", "flex"): 8, ("DCC", "prime"): 44, ("DCC", "rhoso"): 0}
        for site in ("DCB", "DCA", "DCC"):
            for product in product_templates.for_site(site):
                default = [t for g in product["groups"] if g["default"] for t in g["paths"]]
                self.assertEqual(len(default), expected[(site, product["id"])], (site, product["id"]))
                self.assertEqual(len(default), len(set(default)), "default paths must not repeat")
                self.assertEqual((product["mode"], product["immediacy"]), ("regular", "immediate"))

    def test_template_paths_pass_binding_validation(self):
        for site in ("DCB", "DCA", "DCC"):
            for product in product_templates.for_site(site):
                paths = [t for g in product["groups"] for t in g["paths"]]
                if not paths:  # product without a template on this site yet (for example RHOSO on DCA)
                    continue
                rows = [{"tdn": t, "vlan": "1841", "mode": product["mode"], "immediacy": product["immediacy"]} for t in paths]
                self.assertEqual(len(aci.normalize_bindings([], rows)), len(set(paths)), (site, product["id"]))

    def test_unknown_site_and_invalid_file(self):
        self.assertTrue(all(not p["groups"] for p in product_templates.for_site("XYZ")))
        with tempfile.NamedTemporaryFile("w", suffix=".json", delete=False) as f:
            f.write('{"products": {"flex": {"label": "Acme Flex"}}, "sites": {"DCB": {"flex": [{"name": "Rusak", "ports": [{"node": 101, "interfaces": "eth1/9-2"}]}]}}}')
        try:
            with self.assertRaises(ValueError) as ctx:
                product_templates.for_site("DCB", f.name)
            self.assertIn("Rusak", str(ctx.exception))
        finally:
            os.unlink(f.name)


class TemplateStoreTest(unittest.TestCase):
    """The editor's tables -> template groups, and what a save changes (no database, no APIC)."""

    def test_normalize_builds_groups_in_order(self):
        groups = template_store.normalize({
            "groups": [{"name": " Port  ESXi ", "default": True}, {"name": "Uplink", "default": "yes"}],
            "physical": [{"group": 0, "pod": "1", "node": "118", "interface": "Eth1/17"}, {"group": 1, "node": "109", "interface": "eth1/4"}],
            "paths": [{"group": 0, "tdn": "topology/pod-1/protpaths-118-119/pathep-[VPC-25G-CLOUD-A-PG]"}]})
        self.assertEqual(groups, [
            {"name": "Port ESXi", "default": True, "paths": ["topology/pod-1/paths-118/pathep-[eth1/17]",
                                                             "topology/pod-1/protpaths-118-119/pathep-[VPC-25G-CLOUD-A-PG]"]},
            {"name": "Uplink", "default": False, "paths": ["topology/pod-1/paths-109/pathep-[eth1/4]"]}])
        self.assertEqual(template_store.normalize({"groups": []}), [])   # an empty template (no product paths on a site)

    def test_normalize_collects_all_errors(self):
        with self.assertRaises(ValueError) as ctx:
            template_store.normalize({
                "groups": [{"name": "A"}, {"name": "a"}, {"name": ""}, {"name": "Empty"}],
                "physical": [{"group": 0, "node": "x", "interface": "eth1/1"}, {"group": 0, "node": "101", "interface": "eth 1"},
                             {"group": 9, "node": "101", "interface": "eth1/2"}, {"group": True, "node": "101", "interface": "eth1/3"},
                             {"group": 0, "node": "101", "interface": "eth1/5"}],
                "paths": [{"group": 1, "tdn": "uni/bad"}, {"group": 1, "tdn": "topology/pod-1/paths-101/pathep-[eth1/5]"}]})
        message = str(ctx.exception)
        for part in ("Group 2: the name 'a' is used twice", "Group 3: Group name is required", "Physical/PC row 1: Node must be a number",
                     "Physical/PC row 2: Invalid interface / PC", "Physical/PC row 3: choose a group", "Physical/PC row 4: choose a group",
                     "vPC/TDN row 1: Invalid path TDN", "is in the template more than once (groups 'A' and 'Group 2')",
                     "Group 'Empty' has no paths"):
            self.assertIn(part, message)
        for bad in ({}, {"groups": "x"}, {"groups": [{"name": f"G{i}"} for i in range(31)]}):
            with self.assertRaises(ValueError):
                template_store.normalize(bad)

    def test_diff(self):
        a, b, c, d = (f"topology/pod-1/paths-101/pathep-[eth1/{i}]" for i in range(1, 5))
        before = [{"name": "One", "default": True, "paths": [a, b]}, {"name": "Two", "default": False, "paths": [c]}]
        after = [{"name": "Two", "default": True, "paths": [c, b]}, {"name": "Three", "default": False, "paths": [d]}]
        result = template_store.diff(before, after)
        self.assertEqual({i["tdn"]: i["status"] for i in result["items"]}, {a: "delete", b: "moved", c: "same", d: "new"})
        self.assertEqual(next(i for i in result["items"] if i["tdn"] == b)["before_group"], "One")
        self.assertEqual(sorted((g["name"], g["status"]) for g in result["groups"]), [("One", "delete"), ("Three", "new"), ("Two", "change")])
        self.assertEqual(result["summary"], {"new": 1, "delete": 1, "moved": 1, "same": 1, "groups": 3})
        self.assertTrue(result["changed"])
        self.assertEqual(template_store.summary_text(result, before, after), "1 added, 1 removed, 1 moved, 3 group changes (3 → 3 paths)")
        same = template_store.diff(before, before)
        self.assertFalse(same["changed"])
        swapped = template_store.diff(before, list(reversed(before)))
        self.assertTrue(swapped["reordered"] and swapped["changed"])

    def test_permission(self):
        self.assertTrue(template_store.can_edit({"role": "admin", "modules": []}))
        self.assertTrue(template_store.can_edit({"role": "operator", "modules": ["aci", "aci-templates"]}))
        self.assertFalse(template_store.can_edit({"role": "operator", "modules": ["aci"]}))
        self.assertFalse(template_store.can_edit(None))


class IpUsageTest(unittest.TestCase):
    def test_helpers(self):
        import ipaddress
        ips = [ipaddress.ip_address(f"10.253.0.{i}") for i in (2, 3, 4, 7, 9, 10)]
        self.assertEqual(ip_usage.compact(ips), ["10.253.0.2-10.253.0.4", "10.253.0.7", "10.253.0.9-10.253.0.10"])
        self.assertEqual([ip_usage.scope_of(ipaddress.ip_network(n)) for n in ("11.8.158.0/24", "10.252.0.0/16", "100.127.1.0/24", "172.27.0.0/24")],
                         ["public", "private", "private", "private"])
        self.assertEqual(ip_usage.interface("topology/pod-1/protpaths-118-119/pathep-[VPC-25G-CLOUD-A-PG]"), "Pod-1/Node-118-119/VPC-25G-CLOUD-A-PG")
        self.assertEqual(ip_usage.interface("topology/pod-1/paths-101/pathep-[eth1/7]"), "Pod-1/Node-101/eth1/7")


class AuditMigrationTest(unittest.TestCase):
    """The history table of an older ACI Web UI version, copied into the portal database as aci_audit, gets the newer
    columns."""

    def test_old_database_gets_the_apic_user_column(self):
        tmp = tempfile.mkdtemp()
        original = audit.DB_PATH
        audit.DB_PATH = os.path.join(tmp, "old.db")
        try:
            with sqlite3.connect(audit.DB_PATH) as db:   # schema of the version with personal APIC logins
                db.execute("CREATE TABLE aci_audit (id INTEGER PRIMARY KEY AUTOINCREMENT, ts TEXT NOT NULL, finished TEXT,"
                           " username TEXT NOT NULL, site TEXT NOT NULL, apic TEXT NOT NULL, action TEXT NOT NULL,"
                           " target TEXT NOT NULL, product TEXT, status TEXT NOT NULL, summary TEXT NOT NULL,"
                           " details TEXT NOT NULL, error TEXT, client_ip TEXT)")
                db.execute("INSERT INTO aci_audit (ts, username, site, apic, action, target, status, summary, details)"
                           " VALUES ('2026-09-24T10:00:00+07:00', 'aciuser', 'DCB', '172.26.2.250', 'static-port',"
                           " 'uni/tn-A/ap-B/epg-C', 'success', '1 new, 0 changed', '{}')")
            audit.init()
            audit.init()   # running it again changes nothing
            self.assertEqual(audit.get(1)["apic_user"], None)
            record_id = audit.begin("jdoe", "DCB", "172.26.2.250", "static-port", "uni/tn-A/ap-B/epg-C", None,
                                    "1 new, 0 changed", {}, "127.0.0.1", apic_user=SERVICE_USER)
            self.assertEqual((audit.get(record_id)["username"], audit.get(record_id)["apic_user"]), ("jdoe", SERVICE_USER))
            self.assertIn(",Client IP,APIC user", audit.to_csv(audit.all_rows({})).splitlines()[0])
        finally:
            audit.DB_PATH = original
            shutil.rmtree(tmp, ignore_errors=True)

    def test_old_database_gets_the_undo_columns(self):
        tmp = tempfile.mkdtemp()
        original = audit.DB_PATH
        audit.DB_PATH = os.path.join(tmp, "old.db")
        try:
            with sqlite3.connect(audit.DB_PATH) as db:   # schema of the version with the APIC service accounts
                db.execute("CREATE TABLE aci_audit (id INTEGER PRIMARY KEY AUTOINCREMENT, ts TEXT NOT NULL, finished TEXT,"
                           " username TEXT NOT NULL, site TEXT NOT NULL, apic TEXT NOT NULL, action TEXT NOT NULL,"
                           " target TEXT NOT NULL, product TEXT, status TEXT NOT NULL, summary TEXT NOT NULL,"
                           " details TEXT NOT NULL, error TEXT, client_ip TEXT, apic_user TEXT)")
                db.execute("INSERT INTO aci_audit (ts, username, site, apic, action, target, status, summary, details, apic_user)"
                           " VALUES ('2026-09-26T10:00:00+07:00', 'jdoe', 'DCA', '172.24.27.251', 'static-port',"
                           " 'uni/tn-A/ap-B/epg-C', 'success', '1 new, 0 changed', '{}', 'svc_aci')")
            audit.init()
            audit.init()   # running it again changes nothing
            self.assertEqual((audit.get(1)["undo_of"], audit.get(1)["undone_by"]), (None, None))
            # Only a successful undo marks the change it undoes.
            failed = audit.begin("jdoe", "DCA", "172.24.27.251", "static-port", "uni/tn-A/ap-B/epg-C", None,
                                 "undo of change #1: 1 removed, 0 restored", {"undo_of": 1, "items": []}, "127.0.0.1", undo_of=1)
            audit.finish(failed, "failed", "APIC HTTP 400")
            self.assertEqual((audit.get(failed)["undo_of"], audit.get(1)["undone_by"]), (1, None))
            undo = audit.begin("jdoe", "DCA", "172.24.27.251", "static-port", "uni/tn-A/ap-B/epg-C", None,
                               "undo of change #1: 1 removed, 0 restored", {"undo_of": 1, "items": []}, "127.0.0.1", undo_of=1)
            audit.finish(undo, "success")
            self.assertEqual(audit.get(1)["undone_by"], undo)
            rows = list(csv.DictReader(io.StringIO(audit.to_csv(audit.all_rows({})))))
            self.assertEqual({r["ID"]: (r["Undo of"], r["Undone by"]) for r in rows}, {"1": ("", str(undo)), str(failed): ("1", ""), str(undo): ("1", "")})
        finally:
            audit.DB_PATH = original
            shutil.rmtree(tmp, ignore_errors=True)


    def test_portal_database_has_the_aci_table(self):
        with sqlite3.connect(audit.DB_PATH) as db:
            tables = {r[0] for r in db.execute("SELECT name FROM sqlite_master WHERE type = 'table'")}
        self.assertEqual(audit.DB_PATH, os.environ["PORTAL_DB"])
        self.assertIn("aci_audit", tables)
        self.assertNotIn("audit", tables)


class EndToEndTest(unittest.TestCase):
    config_version = 0
    sites_version = 0

    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.mkdtemp(dir=TMP)
        cert, key = make_cert(cls.tmp)
        cls.dcb = MockApic("DCB", cert, key, REFRESH_TIMEOUT).start()
        cls.dca = MockApic("DCA", cert, key, REFRESH_TIMEOUT).start()
        cls.dcc = MockApic("DCC", cert, key, REFRESH_TIMEOUT).start()
        cls.ad_cert, ad_key = fake_ldap.make_cert(cls.tmp)
        cls.ldap = fake_ldap.FakeLdap(cls.ad_cert, ad_key, AD_USERS).start()
        cls.duo = fake_duo.FakeDuo(cls.ad_cert, ad_key).start()
        cls.sso = fake_sso.FakeSso(cls.ad_cert, ad_key, approve_as="operator2").start()
        cls.config_path = os.environ["PORTAL_AUTH"]
        cls.sites_path = os.environ["PORTAL_SITES"]
        cls.write_config()
        cls.write_sites()
        cls.webapp = aci_module
        cls.app = portal.app
        # The administrator from auth.yml adds the operator (with the ACI module) and the service accounts of DCB and
        # DCA. DCC has none.
        admin = cls.sign_in("admin")
        r = admin.post("/api/admin/users", json={"username": "operator2", "role": "operator", "modules": ["aci"]})
        if r.status_code != 200:
            raise RuntimeError(r.get_json())
        # An L1 operator who may only use Operation Automation (a module this portal does not load here): no access to ACI.
        aci_module.PORTAL.store.add_user("operator4", "operator", "test", "wi")
        for site in ("DCB", "DCA"):
            r = admin.post(f"/api/admin/credentials/apic:{site}", json={"username": SERVICE_USER, "password": "secret"})
            if not r.get_json()["test"]["ok"]:
                raise RuntimeError(r.get_json())

    @classmethod
    def tearDownClass(cls):
        for mock in (cls.dcb, cls.dca, cls.dcc, cls.ldap, cls.duo, cls.sso):
            mock.stop()
        shutil.rmtree(cls.tmp, ignore_errors=True)

    @classmethod
    def write_sites(cls, sites=None):
        """sites.yml of the portal (default: DCB, DCA and DCC, each with its mock APIC)."""
        sites = sites or {name: {"apic": {"host": mock.host}} for name, mock in (("DCB", cls.dcb), ("DCA", cls.dca), ("DCC", cls.dcc))}
        with open(cls.sites_path, "w") as f:
            json.dump({"sites": sites}, f)   # JSON is valid YAML
        cls.sites_version += 1   # a new mtime, even within the same second, so the portal reads the file again
        stamp = time.time() + cls.sites_version
        os.utime(cls.sites_path, (stamp, stamp))

    @classmethod
    def write_config(cls, duo=False, sso=False):
        cfg = {"ldap": {"url": [] if sso else [cls.ldap.url], "user_format": "{username}@example.local", "ca_file": cls.ad_cert,
                        "base_dn": "DC=example,DC=local", "timeout_seconds": 5},
               "duo": {"enabled": duo, "client_id": fake_duo.CLIENT_ID, "client_secret": fake_duo.CLIENT_SECRET,
                       "api_host": cls.duo.api_host, "ca_file": cls.ad_cert,
                       "redirect_uri": "https://aciwebui.example.local/auth/duo/callback"},
               "sso": {"enabled": sso, "issuer": cls.sso.issuer, "client_id": fake_sso.CLIENT_ID,
                       "client_secret": fake_sso.CLIENT_SECRET, "ca_file": cls.ad_cert,
                       "redirect_uri": "https://aciwebui.example.local/auth/sso/callback"},
               "admins": ["admin"]}
        with open(cls.config_path, "w") as f:
            json.dump(cfg, f)   # JSON is valid YAML
        cls.config_version += 1   # a new mtime, even within the same second, so the app reloads the file
        stamp = time.time() + cls.config_version
        os.utime(cls.config_path, (stamp, stamp))

    @classmethod
    def sign_in(cls, username="operator2", password="secret"):
        client = cls.app.test_client()
        r = client.post("/api/login", json={"username": username, "password": password})
        if r.status_code != 200 or r.get_json().get("redirect"):
            raise AssertionError(f"sign-in as {username} failed: {r.status_code} {r.get_json()}")
        return client

    def client(self, username="operator2"):
        """Test client signed in with an AD account (operator by default)."""
        return self.sign_in(username)

    def follow_duo(self, url):
        """The browser part: open Duo's authorize URL and return the redirect back to the portal."""
        ctx = ssl.create_default_context(cafile=self.ad_cert)

        class NoRedirect(urllib.request.HTTPRedirectHandler):
            def redirect_request(self, *args, **kwargs):
                return None

        opener = urllib.request.build_opener(NoRedirect, urllib.request.HTTPSHandler(context=ctx), urllib.request.ProxyHandler({}))
        try:
            opener.open(url, timeout=10)
        except urllib.error.HTTPError as e:
            if e.code == 302:
                return e.headers["Location"]
            raise
        self.fail("Duo did not redirect back")

    def make_epg(self, mock, name, app="MGMT-InBand", domains=()):
        """An EPG of tenant CLOUD-1 made directly on a mock APIC; it and everything under it is removed after the test."""
        dn = f"uni/tn-CLOUD-1/ap-{app}/epg-{name}"
        mock.add("fvAEPg", dn, name=name)
        for domain in domains:
            mock.add("fvRsDomAtt", f"{dn}/rsdomAtt-[{domain}]", tDn=domain)
        self.addCleanup(lambda: [mock.mit.pop(k) for k in [k for k in mock.mit if k == dn or k.startswith(dn + "/")]])
        return dn

    @staticmethod
    def bind(mock, epg, tdn, vlan, mode="regular"):
        """A static path made outside the portal, directly on APIC (without modTs, as the mock's own objects)."""
        dn = f"{epg}/rspathAtt-[{tdn}]"
        mock.add("fvRsPathAtt", dn, tDn=tdn, encap=f"vlan-{vlan}", mode=mode, instrImedcy="immediate")
        return dn

    # ---- page, sign-in, module access & session ----
    def test_page_opens_with_the_portal_launcher(self):
        r = self.app.test_client().get("/aci/")
        self.assertEqual(r.status_code, 200)
        self.assertNotIn("WWW-Authenticate", r.headers)
        html = r.data.decode()
        for tab in ("Static Port", "Provisioning", "Report", "History"):
            self.assertIn(tab, html)
        # The portal signs in and runs Users / Credentials; the page has the application launcher of the portal.
        self.assertIn('<div id="portalLauncher" data-current="aci"></div>', html)
        self.assertIn('href="/static/switcher.css"', html)
        self.assertIn('src="/static/switcher.js"', html)
        self.assertLess(html.index('id="portalLauncher"'), html.index('id="userMenu"'))
        for gone in ('id="loginView"', 'id="loginForm"', 'id="tab-users"', 'id="tab-credentials"', 'id="userDialog"', 'id="credDialog"'):
            self.assertNotIn(gone, html)
        self.assertIn('href="/#/users"', html)
        self.assertIn('href="/#/credentials"', html)
        self.assertIn("No access to ACI", html)
        # The page's own files, and the fonts of the portal.
        for path in re.findall(r'(?:href|src)="(/aci/static/[^"]+)"', html):
            self.assertEqual(self.app.test_client().get(path).status_code, 200, path)
        css = self.app.test_client().get("/aci/static/app.css").data.decode()
        fonts = re.findall(r'url\("([^"]+\.woff2)"\)', css)
        self.assertEqual(len(fonts), 6)
        for font in fonts:
            self.assertTrue(font.startswith("/static/fonts/"), font)
            self.assertEqual(self.app.test_client().get(font).status_code, 200, font)
        self.assertEqual(self.app.test_client().get("/aci").status_code, 308)   # /aci -> /aci/

    def test_every_request_of_the_page_goes_to_the_module(self):
        js = self.app.test_client().get("/aci/static/app.js").data.decode()
        self.assertIn("const BASE = '/aci';", js)
        # Module requests go through request(BASE + ...); only the sign-out goes to the portal.
        calls = re.findall(r"request\(([^,)]+)", js)
        self.assertEqual(sorted(set(calls)), ["'/api/logout'", "BASE + '/api/keepalive'", "BASE + '/api/session'", "BASE + url", "url"])
        self.assertIn("PortalLauncher.signIn()", js)
        self.assertIn("location.assign('/')", js)
        self.assertNotIn("/api/login", js)
        self.assertNotIn("/api/admin", js)

    def test_ui_version_is_announced(self):
        # A browser tab that still runs an older version can tell that the server was updated.
        html = self.app.test_client().get("/aci/").data.decode()
        version = re.search(r'<meta name="ui-version" content="([0-9a-f]{12})">', html).group(1)
        self.assertEqual(version, aci_module.UI_VERSION)
        self.assertIn(f"/aci/static/app.js?v={version}", html)
        self.assertIn(f"/aci/static/app.css?v={version}", html)
        self.assertEqual(self.app.test_client().get("/aci/api/session").get_json()["ui_version"], version)
        self.assertEqual(self.client().get("/aci/api/session").get_json()["ui_version"], version)
        self.assertEqual(self.client().post("/aci/api/keepalive").get_json()["ui_version"], version)
        r = self.app.test_client().get("/aci/api/discover?kind=tenants&site=DCB")
        self.assertEqual((r.status_code, r.get_json()["session"]["ui_version"]), (401, version))

    def test_session_shows_user_role_access_and_service_accounts(self):
        anonymous = self.app.test_client().get("/aci/api/session").get_json()
        self.assertEqual((anonymous["ok"], anonymous["user"], anonymous["allowed"]), (True, None, False))
        self.assertEqual(anonymous["auth"], {"method": "Active Directory", "duo": False, "local_accounts": False, "sso": False, "sso_label": "Sign in with Duo SSO", "ad_password": True, "ad_push": False, "problems": []})
        # Without a session only the site names and hosts are shown, not the state of the service accounts.
        self.assertEqual(anonymous["sites"], [{"name": "DCB", "host": self.dcb.host}, {"name": "DCA", "host": self.dca.host},
                                              {"name": "DCC", "host": self.dcc.host}])
        status = self.client().get("/aci/api/session").get_json()
        self.assertEqual(status["user"], {"username": "operator2", "display_name": "Otto Operator", "role": "operator", "mfa": "none",
                                          "source": "ad", "modules": ["aci"]})
        self.assertTrue(status["allowed"])
        self.assertEqual([(s["name"], s["host"], s["configured"], s["error"]) for s in status["sites"]],
                         [("DCB", self.dcb.host, True, ""), ("DCA", self.dca.host, True, ""), ("DCC", self.dcc.host, False, "")])
        admin = self.client("admin").get("/aci/api/session").get_json()
        self.assertEqual((admin["user"]["username"], admin["user"]["display_name"], admin["user"]["role"], admin["allowed"]),
                         ("admin", "Ada Admin", "admin", True))
        # The portal's own session lists the module.
        modules = self.client().get("/api/session").get_json()["modules"]
        self.assertEqual([(m["id"], m["title"], m["url"], m["icon"], m["allowed"]) for m in modules],
                         [("aci", "ACI", "/aci/", "network", True)])

    def test_operator_without_the_aci_module(self):
        client = self.client("operator4")   # signed in on the portal, but ACI is not one of the account's modules
        status = client.get("/aci/api/session")
        self.assertEqual(status.status_code, 200)   # never 403: the page shows "No access to ACI" itself
        status = status.get_json()
        self.assertEqual((status["user"]["username"], status["user"]["modules"], status["allowed"]), ("operator4", [], False))
        self.assertEqual(status["sites"], [{"name": "DCB", "host": self.dcb.host}, {"name": "DCA", "host": self.dca.host},
                                           {"name": "DCC", "host": self.dcc.host}])
        requests = len(self.dcb.requests)
        r = client.get("/aci/api/discover?kind=tenants&site=DCB")
        self.assertEqual(r.status_code, 403)
        self.assertIn("Your account has no access to ACI", r.get_json()["error"])
        for method, url, body in (("get", "/aci/api/report?site=DCB", None), ("get", "/aci/api/audit", None),
                                  ("get", "/aci/api/audit/1", None), ("get", "/aci/export/audit.csv", None),
                                  ("get", "/aci/export/report.csv?site=DCB", None),
                                  ("post", "/aci/api/static-port/preview?site=DCB", {"tenant": "CLOUD-1"}),
                                  ("post", "/aci/api/provision/apply?site=DCB", {}),
                                  ("post", "/aci/api/audit/1/undo", {"confirm": True})):
            with self.subTest(url=url):
                kwargs = {"json": body} if body is not None else {}
                self.assertEqual(getattr(client, method)(url, **kwargs).status_code, 403)
        self.assertEqual(len(self.dcb.requests), requests, "APIC is never contacted for an account without ACI")
        # keepalive keeps the portal sign-in alive and reports the same, without touching APIC.
        r = client.post("/aci/api/keepalive")
        self.assertEqual((r.status_code, r.get_json()["user"]["username"], r.get_json()["allowed"]), (200, "operator4", False))
        self.assertEqual(self.app.test_client().get("/aci/api/discover?kind=tenants&site=DCB").status_code, 401)
        # The page itself opens (with the launcher), and the portal lists ACI as not allowed for this account.
        self.assertIn('data-current="aci"', client.get("/aci/").data.decode())
        modules = client.get("/api/session").get_json()["modules"]
        self.assertEqual([(m["id"], m["allowed"]) for m in modules], [("aci", False)])
        # Once an administrator gives the account ACI, it works at once.
        admin = self.client("admin")
        self.assertEqual(admin.post("/api/admin/users/operator4", json={"modules": ["aci"]}).status_code, 200)
        try:
            self.assertTrue(client.get("/aci/api/session").get_json()["allowed"])
            self.assertEqual(client.get("/aci/api/discover?kind=tenants&site=DCB").status_code, 200)
        finally:
            aci_module.PORTAL.store.update_user("operator4", modules="wi")
        self.assertEqual(client.get("/aci/api/discover?kind=tenants&site=DCB").status_code, 403)

    def test_sign_in_is_refused_for_outsiders_and_wrong_passwords(self):
        client = self.app.test_client()
        r = client.post("/api/login", json={"username": "outsider", "password": "secret"})
        self.assertEqual(r.status_code, 403)
        self.assertIn("not allowed to use Ops Portal", r.get_json()["error"])
        r = client.post("/api/login", json={"username": "operator2", "password": "wrong"})
        self.assertEqual((r.status_code, r.get_json()["error"]), (401, "Invalid username or password."))
        self.assertIsNone(client.get("/aci/api/session").get_json()["user"])
        self.assertEqual(client.get("/aci/api/discover?kind=tenants&site=DCB").status_code, 401)
        # "EXAMPLE\\jdoe" and "jdoe@example.local" are the same account.
        self.assertIsNotNone(self.sign_in("EXAMPLE\\Operator2").get("/aci/api/session").get_json()["user"])

    def test_duo_sign_in_through_the_app(self):
        self.write_config(duo=True)
        try:
            client = self.app.test_client()
            r = client.post("/api/login", json={"username": "operator2", "password": "secret"})
            self.assertEqual(r.status_code, 200, r.get_json())
            url = r.get_json()["redirect"]
            self.assertTrue(url.startswith(f"https://{self.duo.api_host}/oauth/v1/authorize?"))
            status = client.get("/aci/api/session").get_json()
            self.assertEqual((status["user"], status["auth"]["duo"]), (None, True))   # not signed in before Duo approves
            self.assertEqual(client.get("/aci/api/discover?kind=tenants&site=DCB").status_code, 401)
            back = urllib.parse.urlsplit(self.follow_duo(url))
            self.assertEqual(back.path, "/auth/duo/callback")
            r = client.get(f"/auth/duo/callback?{back.query}")
            location = urllib.parse.urlsplit(r.headers["Location"])
            self.assertEqual((r.status_code, location.path, location.query), (302, "/", ""))
            user = client.get("/aci/api/session").get_json()["user"]
            self.assertEqual((user["username"], user["role"], user["mfa"]), ("operator2", "operator", "duo"))
            self.assertEqual(client.get("/aci/api/discover?kind=tenants&site=DCB").status_code, 200)
            # The same Duo answer cannot be used again: the browser comes back with an error for the login page.
            replay = self.app.test_client().get(f"/auth/duo/callback?{back.query}")
            self.assertIn("auth_error=", replay.headers["Location"])
        finally:
            self.write_config(duo=False)

    def test_duo_sso_sign_in_through_the_app(self):
        self.write_config(sso=True)
        try:
            client = self.app.test_client()
            auth = client.get("/aci/api/session").get_json()["auth"]
            self.assertEqual((auth["sso"], auth["ad_password"]), (True, False))
            r = client.post("/api/login", json={"username": "operator2", "password": "secret"})   # AD accounts: SSO only
            self.assertEqual(r.status_code, 400)
            self.assertIn("Duo SSO", r.get_json()["error"])
            r = client.get("/auth/sso/login")   # opened by another name: first moved to the host of redirect_uri
            self.assertEqual((r.status_code, r.headers["Location"]), (302, "https://aciwebui.example.local/auth/sso/login"))
            r = client.get("/auth/sso/login", base_url="https://aciwebui.example.local")
            self.assertEqual(r.status_code, 302)
            back = urllib.parse.urlsplit(self.follow_duo(r.headers["Location"]))
            self.assertEqual(back.path, "/auth/sso/callback")
            r = client.get(f"/auth/sso/callback?{back.query}", base_url="https://aciwebui.example.local")
            location = urllib.parse.urlsplit(r.headers["Location"])
            self.assertEqual((r.status_code, location.path, location.query), (302, "/", ""))
            user = client.get("/aci/api/session", base_url="https://aciwebui.example.local").get_json()["user"]   # the session cookie is per host
            self.assertEqual((user["username"], user["role"], user["mfa"]), ("operator2", "operator", "sso"))
            self.assertEqual(client.get("/aci/api/discover?kind=tenants&site=DCB", base_url="https://aciwebui.example.local").status_code, 200)
        finally:
            self.write_config()

    def test_logout(self):
        client = self.client()
        self.assertEqual(client.get("/aci/api/discover?kind=tenants&site=DCB").status_code, 200)
        before = self.dcb.count("POST", "/api/aaaLogout.json")
        self.assertEqual(client.post("/api/logout").status_code, 200)
        self.assertIsNone(client.get("/aci/api/session").get_json()["user"])
        self.assertEqual(client.get("/aci/api/discover?kind=tenants&site=DCB").status_code, 401)
        # The APIC token belongs to the shared service account: signing out of the portal does not log it out.
        self.assertEqual(self.dcb.count("POST", "/api/aaaLogout.json"), before)
        self.assertEqual(self.client().get("/aci/api/discover?kind=tenants&site=DCB").status_code, 200)

    # ---- service accounts ----
    def test_site_is_validated_per_request(self):
        client = self.client()
        self.assertEqual(client.get("/aci/api/discover?kind=tenants&site=EVIL").status_code, 400)
        self.assertEqual(client.get("/aci/api/discover?kind=tenants").status_code, 400)
        requests = len(self.dcc.requests)
        r = client.get("/aci/api/discover?kind=tenants&site=DCC")   # DCC has no service account
        self.assertEqual(r.status_code, 400)
        self.assertIn("The APIC DCC service account is not set", r.get_json()["error"])
        self.assertEqual(len(self.dcc.requests), requests, "without a service account APIC is not contacted")
        self.assertEqual(self.app.test_client().get("/aci/api/discover?kind=tenants&site=DCB").status_code, 401)
        self.assertEqual(client.get("/aci/api/discover?kind=apps&tenant=a/b&site=DCB").status_code, 400)

    def test_service_account_changes_take_effect_at_once(self):
        admin, operator = self.client("admin"), self.client()
        listing = admin.get("/api/admin/credentials").get_json()["credentials"]
        self.assertEqual([(c["name"], c["label"], c["host"], c["site"], c["configured"], c["username"]) for c in listing],
                         [("apic:DCB", "APIC DCB", self.dcb.host, "DCB", True, SERVICE_USER),
                          ("apic:DCA", "APIC DCA", self.dca.host, "DCA", True, SERVICE_USER),
                          ("apic:DCC", "APIC DCC", self.dcc.host, "DCC", False, "")])
        self.assertEqual({(c["kind"], c["module"], c["module_title"]) for c in listing}, {("apic", "aci", "ACI")})
        try:
            # A wrong password is saved, but the connection test and every use of the site explain the problem.
            r = admin.post("/api/admin/credentials/apic:DCC", json={"username": SERVICE_USER, "password": "wrong"})
            test = r.get_json()["test"]
            self.assertFalse(test["ok"])
            self.assertIn(f"Cannot log in to APIC {self.dcc.host} as {SERVICE_USER}: APIC HTTP 401", test["message"])
            r = operator.get("/aci/api/discover?kind=tenants&site=DCC")
            self.assertEqual(r.status_code, 502)
            self.assertIn(f"APIC DCC ({self.dcc.host}) refused the service account {SERVICE_USER}: APIC HTTP 401", r.get_json()["error"])
            self.assertIn("Administration → Credentials", r.get_json()["error"])
            dcc = next(s for s in operator.get("/aci/api/session").get_json()["sites"] if s["name"] == "DCC")
            self.assertTrue(dcc["configured"])
            self.assertIn("refused the service account", dcc["error"])
            # The right password: the old state is dropped and DCC works at once, without a restart.
            r = admin.post("/api/admin/credentials/apic:DCC", json={"username": SERVICE_USER, "password": "secret"})
            self.assertEqual(r.get_json()["test"], {"ok": True, "message": f"Logged in to APIC {self.dcc.host} as {SERVICE_USER}"})
            self.assertEqual(admin.post("/api/admin/credentials/apic:DCC/test", json={}).get_json()["test"]["ok"], True)
            self.assertIn("DCC-Only", operator.get("/aci/api/discover?kind=tenants&site=DCC").get_json()["items"])
            dcc = next(s for s in operator.get("/aci/api/session").get_json()["sites"] if s["name"] == "DCC")
            self.assertEqual((dcc["configured"], dcc["error"]), (True, ""))
            listing = admin.get("/api/admin/credentials").get_json()
            self.assertNotIn('"secret"', json.dumps(listing))   # the password is never shown again
            self.assertEqual(next(c for c in listing["credentials"] if c["site"] == "DCC")["updated_by"], "admin")
        finally:
            self.assertEqual(admin.delete("/api/admin/credentials/apic:DCC").status_code, 200)
        r = operator.get("/aci/api/discover?kind=tenants&site=DCC")
        self.assertEqual(r.status_code, 400)
        self.assertIn("service account is not set", r.get_json()["error"])
        events = [(e["actor"], e["action"], e["target"]) for e in admin.get("/api/admin/events?limit=20").get_json()["events"]]
        self.assertIn(("admin", "credential.set", "apic:DCC"), events)
        self.assertIn(("admin", "credential.remove", "apic:DCC"), events)

    def test_apic_is_reached_directly_even_with_proxy_env(self):
        # A dead proxy on port 9: if the app used the proxy variables, the login to APIC would fail.
        client = self.client()
        self.webapp.CONNECTIONS.drop("DCB")   # the next call logs in to APIC again
        logins = self.dcb.count("POST", "/api/aaaLogin.json")
        proxy_env = {name: "http://127.0.0.1:9" for name in ("http_proxy", "https_proxy", "HTTP_PROXY", "HTTPS_PROXY")}
        with unittest.mock.patch.dict(os.environ, proxy_env):
            for name in ("no_proxy", "NO_PROXY"):
                os.environ.pop(name, None)
            r = client.get("/aci/api/discover?kind=tenants&site=DCB")
        self.assertEqual(r.status_code, 200, r.get_json())
        self.assertEqual(self.dcb.count("POST", "/api/aaaLogin.json"), logins + 1)

    def test_refused_token_logs_in_again_once(self):
        client = self.client()
        self.webapp.CONNECTIONS.drop("DCB")
        self.assertEqual(client.get("/aci/api/discover?kind=tenants&site=DCB").status_code, 200)   # fresh token
        logins = self.dcb.count("POST", "/api/aaaLogin.json")
        with self.dcb.lock:
            self.dcb.tokens.clear()   # for example after an APIC restart: the token is refused with HTTP 403
        r = client.get("/aci/api/discover?kind=tenants&site=DCB")
        self.assertEqual(r.status_code, 200, r.get_json())
        self.assertEqual(self.dcb.count("POST", "/api/aaaLogin.json"), logins + 1)

    def test_sites_come_from_sites_yml_and_are_read_again(self):
        client, admin = self.client(), self.client("admin")
        self.assertIn("DCB-Only", client.get("/aci/api/discover?kind=tenants&site=DCB").get_json()["items"])
        dcb_logins, dca_logins = self.dcb.count("POST", "/api/aaaLogin.json"), self.dca.count("POST", "/api/aaaLogin.json")
        # sites.yml changes while the portal runs: DCB gets another APIC (here the mock of DCA), site XYZ is added, DCC
        # is removed, and a site without `apic` (used by other modules only) is no ACI site.
        try:
            self.write_sites({"DCB": {"apic": {"host": self.dca.host}}, "DCA": {"apic": {"host": self.dca.host}},
                              "XYZ": {"apic": {"host": self.dcc.host}}, "NOACI": {"ontap": {"host": "10.253.0.1"}}})
            status = client.get("/aci/api/session").get_json()
            self.assertEqual([(s["name"], s["host"], s["configured"]) for s in status["sites"]],
                             [("DCB", self.dca.host, True), ("DCA", self.dca.host, True), ("XYZ", self.dcc.host, False)])
            targets = [(c["name"], c["host"]) for c in admin.get("/api/admin/credentials").get_json()["credentials"]]
            self.assertEqual(targets, [("apic:DCB", self.dca.host), ("apic:DCA", self.dca.host), ("apic:XYZ", self.dcc.host)])
            # The token of the old DCB APIC is not used for the new one: the next call logs in there.
            tenants = client.get("/aci/api/discover?kind=tenants&site=DCB").get_json()["items"]
            self.assertIn("DCA-Only", tenants)
            self.assertNotIn("DCB-Only", tenants)
            self.assertEqual(self.dcb.count("POST", "/api/aaaLogin.json"), dcb_logins)
            self.assertEqual(self.dca.count("POST", "/api/aaaLogin.json"), dca_logins + 1)
            for site in ("DCC", "NOACI"):
                r = client.get(f"/aci/api/discover?kind=tenants&site={site}")
                self.assertEqual((r.status_code, r.get_json()["error"]), (400, f"Unknown site: {site}"))
            r = client.get("/aci/api/discover?kind=tenants&site=XYZ")
            self.assertEqual(r.status_code, 400)
            self.assertIn("The APIC XYZ service account is not set", r.get_json()["error"])
        finally:
            self.write_sites()
        self.assertIn("DCB-Only", client.get("/aci/api/discover?kind=tenants&site=DCB").get_json()["items"])
        self.assertEqual([s["name"] for s in client.get("/aci/api/session").get_json()["sites"]], ["DCB", "DCA", "DCC"])

    # ---- discovery ----
    def test_discovery_is_per_site(self):
        client = self.client()
        self.assertIn("DCB-Only", client.get("/aci/api/discover?kind=tenants&site=DCB").get_json()["items"])
        self.assertIn("DCA-Only", client.get("/aci/api/discover?kind=tenants&site=DCA").get_json()["items"])
        domains = client.get("/aci/api/discover?kind=domains&site=DCA").get_json()["items"]
        self.assertIn({"dn": "uni/vmmp-OpenStack/dom-DCA-OSP", "label": "OpenStack / DCA-OSP"}, domains)
        self.assertIn({"dn": "uni/vmmp-VMware/dom-DCA-DVS", "label": "VMware / DCA-DVS"}, domains)
        objects = client.get("/aci/api/discover?kind=tenant_objects&tenant=CLOUD-1&site=DCB").get_json()["items"]
        self.assertEqual(objects, {"vrfs": ["VRF-1"], "bds": ["BD-1", "BD-2"], "apps": ["MGMT-InBand", "PUBLIC-CLOUD-OPENSTACK"], "l3outs": ["L3OUT-1"]})

    def test_path_search(self):
        client = self.client()
        physical = [p["dn"] for p in client.get("/aci/api/paths?kind=physical&q=101&site=DCA").get_json()["items"]]
        self.assertIn("topology/pod-1/paths-101/pathep-[PC-SERVER-01]", physical)
        self.assertIn("topology/pod-1/paths-101/extpaths-111/pathep-[eth1/5]", physical)
        self.assertFalse(any("protpaths" in dn for dn in physical))
        vpc = [p["dn"] for p in client.get("/aci/api/paths?kind=vpc&site=DCA").get_json()["items"]]
        self.assertIn("topology/pod-1/protpaths-103-104/pathep-[Openstack-Traffic-A-PG]", vpc)
        self.assertTrue(vpc and all("/protpaths-" in dn for dn in vpc))

    # ---- static port ----
    def test_static_port_applies_only_to_selected_site(self):
        client = self.client()
        req = {"tenant": "CLOUD-1", "app": "MGMT-InBand", "epg": "MGMT-VYOS",
               "physical": [{"pod": "1", "node": "106", "interface": "eth1/17", "vlan": "25", "mode": "regular", "immediacy": "immediate"}],
               "paths": [{"tdn": "topology/pod-1/protpaths-103-104/pathep-[Openstack-Traffic-A-PG]", "vlan": "25", "mode": "regular", "immediacy": "immediate"}]}
        dcb_posts, dca_posts = len(self.dcb.posts), len(self.dca.posts)
        preview = client.post("/aci/api/static-port/preview?site=DCA", json=req).get_json()
        self.assertEqual(preview["summary"], {"new": 2, "change": 0, "same": 0})
        self.assertEqual(len(self.dca.posts), dca_posts, "a preview must not change APIC")
        result = client.post("/aci/api/static-port/apply?site=DCA", json=req).get_json()
        self.assertTrue(result["applied"])
        self.assertEqual(len(self.dca.posts), dca_posts + 1, "all bindings are sent in one POST")
        dn = "uni/tn-CLOUD-1/ap-MGMT-InBand/epg-MGMT-VYOS/rspathAtt-[topology/pod-1/paths-106/pathep-[eth1/17]]"
        self.assertEqual(self.dca.mit[dn][1]["encap"], "vlan-25")
        self.assertNotIn(dn, self.dcb.mit)
        self.assertEqual(len(self.dcb.posts), dcb_posts)
        again = client.post("/aci/api/static-port/apply?site=DCA", json=req).get_json()
        self.assertEqual(again["summary"], {"new": 0, "change": 0, "same": 2})
        self.assertEqual(len(self.dca.posts), dca_posts + 1, "an identical binding is not sent again")

    def test_static_port_preview_shows_changed_vlan(self):
        client = self.client()
        req = {"tenant": "CLOUD-1", "app": "MGMT-InBand", "epg": "Mgmt-InBand-EPG", "physical": [{"node": "101", "interface": "eth1/1", "vlan": "20"}]}
        preview = client.post("/aci/api/static-port/preview?site=DCA", json=req).get_json()
        self.assertEqual(preview["summary"], {"new": 0, "change": 1, "same": 0})
        self.assertEqual(preview["items"][0]["before"]["encap"], "vlan-11")

    def test_static_port_errors(self):
        client = self.client()
        r = client.post("/aci/api/static-port/preview?site=DCA", json={"tenant": "CLOUD-1", "app": "MGMT-InBand", "epg": "DOES-NOT-EXIST",
                                                                  "physical": [{"node": "101", "interface": "eth1/9", "vlan": "10"}]})
        self.assertEqual(r.status_code, 400)
        self.assertIn("not found", r.get_json()["error"])
        r = client.post("/aci/api/static-port/preview?site=DCA", json={"tenant": "CLOUD-1", "app": "MGMT-InBand", "epg": "MGMT-VYOS",
                                                                  "physical": [{"node": "101", "interface": "eth1/9", "vlan": "9999"}]})
        self.assertEqual(r.status_code, 400)
        self.assertIn("VLAN must be 1-4094", r.get_json()["error"])

    def test_vlan_used_by_another_epg_is_refused_while_typing_and_at_apply(self):
        # VLAN 300 on a static path and VLAN 301 on a static leaf of EPG CLOUD-1/PUBLIC-CLOUD-OPENSTACK/MGMT-VYOS,
        # VLAN 302 on an interface of L3Out CLOUD-1/L3OUT-1 (leaf 101).
        added = self.dca.add_vlan_users()
        self.addCleanup(lambda: [self.dca.mit.pop(dn, None) for dn in added])
        client = self.client()
        epg = {"tenant": "CLOUD-1", "app": "MGMT-InBand", "epg": "Mgmt-InBand-EPG"}
        # 1) While typing: the live check tells where each VLAN is used (the EPG's own VLAN 11 is free for it).
        r = client.get("/aci/api/vlan-usage?site=DCA&" + urllib.parse.urlencode(dict(epg, vlans="300,301,302,11,999")))
        self.assertEqual(r.status_code, 200, r.get_json())
        usage = r.get_json()["usage"]
        self.assertEqual({v: [(u["kind"], u["tenant"], u["app"], u["name"], u["binding"]) for u in usage[v]] for v in usage},
                         {"300": [("EPG", "CLOUD-1", "PUBLIC-CLOUD-OPENSTACK", "MGMT-VYOS", "static path")],
                          "301": [("EPG", "CLOUD-1", "PUBLIC-CLOUD-OPENSTACK", "MGMT-VYOS", "static leaf")],
                          "302": [("L3Out", "CLOUD-1", "", "L3OUT-1", "L3Out interface")], "11": [], "999": []})
        self.assertEqual(usage["300"][0]["nodes"], [106])
        # Without an EPG every user is listed: VLAN 11 belongs to Mgmt-InBand-EPG.
        other = client.get("/aci/api/vlan-usage?site=DCA&vlans=11").get_json()["usage"]["11"]
        self.assertEqual({u["name"] for u in other}, {"Mgmt-InBand-EPG"})
        # The EPG that owns VLAN 300 may use it on more ports.
        own = client.get("/aci/api/vlan-usage?site=DCA&tenant=CLOUD-1&app=PUBLIC-CLOUD-OPENSTACK&epg=MGMT-VYOS&vlans=300").get_json()
        self.assertEqual(own["usage"]["300"], [])
        for bad, message in (("", "No VLAN to check"), ("5000", "VLAN must be 1-4094"), ("x", "VLAN must be a number")):
            r = client.get(f"/aci/api/vlan-usage?site=DCA&vlans={bad}")
            self.assertEqual(r.status_code, 400)
            self.assertIn(message, r.get_json()["error"])
        # 2) Review: the same check on the plan; nothing can be applied.
        req = dict(epg, physical=[{"node": "106", "interface": "eth1/20", "vlan": "300"},
                                   {"node": "106", "interface": "eth1/14", "vlan": "302"},
                                   {"node": "101", "interface": "eth1/15", "vlan": "11"}])
        posts = len(self.dca.posts)
        preview = client.post("/aci/api/static-port/preview?site=DCA", json=req).get_json()
        self.assertTrue(preview["blocked"])
        self.assertIsNone(preview["payload"])
        self.assertEqual([c["vlan"] for c in preview["conflicts"]], [300, 302])
        self.assertTrue(preview["conflicts"][0]["used_by"][0]["same_leaf"])    # also on leaf 106
        self.assertFalse(preview["conflicts"][1]["used_by"][0]["same_leaf"])   # the L3Out is on leaf 101: still refused
        self.assertEqual([i.get("conflict", False) for i in preview["items"]], [True, True, False])
        # 3) Apply: checked again on the server; refused, nothing sent, nothing recorded.
        audits = self.webapp.audit.query({}, 1000, 0)[0]
        r = client.post("/aci/api/static-port/apply?site=DCA", json=req)
        self.assertEqual(r.status_code, 400)
        error = r.get_json()["error"]
        self.assertIn("VLAN 300 is already used by EPG CLOUD-1/PUBLIC-CLOUD-OPENSTACK/MGMT-VYOS", error)
        self.assertIn("VLAN 302 is already used by L3Out CLOUD-1/L3OUT-1", error)
        self.assertIn("Use a new VLAN ID", error)
        self.assertEqual(len(self.dca.posts), posts)
        self.assertEqual(self.webapp.audit.query({}, 1000, 0)[0], audits)
        # 4) Changing only the mode of an existing binding keeps its VLAN: not blocked, even if another EPG uses it too.
        extra = "uni/tn-CLOUD-1/ap-PUBLIC-CLOUD-OPENSTACK/epg-Mgmt-InBand-EPG/rspathAtt-[topology/pod-1/paths-106/pathep-[eth1/30]]"
        self.dca.add("fvRsPathAtt", extra, tDn="topology/pod-1/paths-106/pathep-[eth1/30]", encap="vlan-11", mode="regular", instrImedcy="immediate")
        try:
            mode_only = dict(epg, physical=[{"node": "101", "interface": "eth1/1", "vlan": "11", "mode": "native"}])
            preview = client.post("/aci/api/static-port/preview?site=DCA", json=mode_only).get_json()
            self.assertEqual((preview["summary"]["change"], preview["blocked"]), (1, False))
            new_port = dict(epg, physical=[{"node": "101", "interface": "eth1/16", "vlan": "11"}])
            preview = client.post("/aci/api/static-port/preview?site=DCA", json=new_port).get_json()
            self.assertTrue(preview["blocked"])   # a new port with VLAN 11: now also used by another EPG
        finally:
            self.dca.mit.pop(extra)

    def test_free_vlans_are_suggested_from_the_static_pool_ranges(self):
        client = self.client()
        epg = {"tenant": "CLOUD-1", "app": "MGMT-InBand", "epg": "Mgmt-InBand-EPG"}
        epg_dn = "uni/tn-CLOUD-1/ap-MGMT-InBand/epg-Mgmt-InBand-EPG"
        posts = len(self.dca.posts)
        # The EPG has a physical domain (static pool: 300-310 plus a dynamic block 1900-1909 that does not count), a
        # VMware domain (dynamic pool with one static block 320-322) and an OpenStack domain without a pool.
        d = client.get("/aci/api/vlan-suggest", query_string={"site": "DCA", **epg}).get_json()
        self.assertTrue(d["ok"], d)
        self.assertEqual([(p["name"], p["domain"], p["ranges"], p["free_count"]) for p in d["pools"]],
                         [("DCA-PHYS-VLANS", "Physical / DCA-PHYS", "300-310", 11), ("DCA-DVS-VLANS", "VMware / DCA-DVS", "320-322", 3)])
        self.assertEqual((d["site"], d["epg_dn"], d["own"], d["free"], d["free_total"], d["note"]),
                         ("DCA", epg_dn, [11], list(range(300, 310)), 14, ""))   # the first 10; VLAN 11 is the EPG's own
        # Used VLANs are never suggested: 300 (static path of another EPG), 301 (static leaf), 302 (L3Out), and 305 of the
        # EPG itself, which is one of its own VLANs now.
        added = self.dca.add_vlan_users()
        added.append(self.bind(self.dca, epg_dn, "topology/pod-1/paths-101/pathep-[eth1/44]", 305))
        self.addCleanup(lambda: [self.dca.mit.pop(dn, None) for dn in added])
        d = client.get("/aci/api/vlan-suggest", query_string={"site": "DCA", "count": 5, **epg}).get_json()
        self.assertEqual((d["own"], d["free"], d["free_total"], [p["free_count"] for p in d["pools"]]),
                         ([11, 305], [303, 304, 306, 307, 308], 10, [7, 3]))
        # An EPG without a domain, and an EPG whose domains have no pool or only dynamic blocks: nothing to suggest.
        self.make_epg(self.dca, "NO-DOMAIN")
        d = client.get("/aci/api/vlan-suggest", query_string={"site": "DCA", **epg, "epg": "NO-DOMAIN"}).get_json()
        self.assertEqual((d["own"], d["free"], d["pools"]), ([], [], []))
        self.assertIn("has no physical or VMM domain", d["note"])
        pool = "uni/infra/vlanns-[DYN-ONLY]-dynamic"
        extra = {"fvnsVlanInstP": (pool, {"name": "DYN-ONLY", "allocMode": "dynamic"}),
                 "fvnsEncapBlk": (f"{pool}/from-[vlan-2500]-to-[vlan-2599]", {"from": "vlan-2500", "to": "vlan-2599", "allocMode": "inherit"}),
                 "vmmDomP": ("uni/vmmp-VMware/dom-DYN", {"name": "DYN"}),
                 "infraRsVlanNs": ("uni/vmmp-VMware/dom-DYN/rsvlanNs", {"tDn": pool})}
        for cls, (dn, attrs) in extra.items():
            self.dca.add(cls, dn, **attrs)
            self.addCleanup(self.dca.mit.pop, dn)
        self.make_epg(self.dca, "DYN-ONLY", domains=["uni/vmmp-VMware/dom-DYN", "uni/vmmp-OpenStack/dom-DCA-OSP"])
        d = client.get("/aci/api/vlan-suggest", query_string={"site": "DCA", **epg, "epg": "DYN-ONLY"}).get_json()
        self.assertEqual((d["free"], d["pools"]), ([], []))
        self.assertIn("no VLAN pool on domain OpenStack / DCA-OSP", d["note"])
        self.assertIn("only dynamic VLAN blocks in pool DYN-ONLY", d["note"])
        # Input is validated like the other APIs; the suggestion only reads APIC.
        for params, message in (({"epg": ""}, "EPG is invalid"), ({"epg": "NOT-THERE"}, "not found on this site"),
                                ({"count": "0"}, "Count must be greater than 0"), ({"count": "101"}, "Count must be 1-100"),
                                ({"count": "x"}, "Count must be a number")):
            r = client.get("/aci/api/vlan-suggest", query_string={"site": "DCA", **epg, **params})
            self.assertEqual(r.status_code, 400, params)
            self.assertIn(message, r.get_json()["error"])
        self.assertEqual(len(self.dca.posts), posts)
        self.assertEqual(self.app.test_client().get("/aci/api/vlan-suggest", query_string={"site": "DCA", **epg}).status_code, 401)

    def test_templates_endpoint_flags_missing_paths(self):
        client = self.client()
        products = client.get("/aci/api/templates?site=DCA").get_json()["products"]
        prime = next(p for p in products if p["id"] == "prime")
        paths = [x for g in prime["groups"] for x in g["paths"]]
        self.assertEqual(len(paths), 48)
        self.assertTrue(all(x["exists"] for x in paths))
        tdn = "topology/pod-1/paths-111/pathep-[eth1/1]"
        removed = self.dca.mit.pop(tdn)
        try:
            products = client.get("/aci/api/templates?site=DCA").get_json()["products"]
            flags = {x["tdn"]: x["exists"] for p in products for g in p["groups"] for x in g["paths"]}
            self.assertFalse(flags[tdn])
        finally:
            self.dca.mit[tdn] = removed
        self.assertEqual(client.get("/aci/api/templates?site=EVIL").status_code, 400)

    def test_apply_full_template_in_one_post(self):
        client = self.client()
        prime = next(p for p in client.get("/aci/api/templates?site=DCA").get_json()["products"] if p["id"] == "prime")
        paths = [x["tdn"] for g in prime["groups"] if g["default"] for x in g["paths"]]
        req = {"tenant": "CLOUD-1", "app": "PUBLIC-CLOUD-OPENSTACK", "epg": "Mgmt-InBand-EPG", "physical": [],
               "paths": [{"tdn": t, "vlan": "1210", "mode": "regular", "immediacy": "immediate"} for t in paths]}
        posts = len(self.dca.posts)
        result = client.post("/aci/api/static-port/apply?site=DCA", json=req).get_json()
        self.assertTrue(result["ok"], result)
        self.assertEqual(result["summary"]["new"], 46)
        self.assertEqual(len(self.dca.posts), posts + 1)
        prefix = "uni/tn-CLOUD-1/ap-PUBLIC-CLOUD-OPENSTACK/epg-Mgmt-InBand-EPG/rspathAtt-"
        self.assertEqual(sum(1 for dn in self.dca.mit if dn.startswith(prefix)), 46)

    # ---- product templates (Fabric -> Templates) ----
    @staticmethod
    def editor_body(product):
        """What the template editor sends for a product of /aci/api/templates, unchanged: groups, plus its rows."""
        body = {"version": product["version"], "groups": [{"name": g["name"], "default": g["default"]} for g in product["groups"]],
                "physical": [], "paths": []}
        for i, g in enumerate(product["groups"]):
            for x in g["paths"]:
                m = re.match(r"^topology/pod-(\d+)/paths-(\d+)/pathep-\[([^\[\]]+)\]$", x["tdn"])
                if m:
                    body["physical"].append({"group": i, "pod": m[1], "node": m[2], "interface": m[3]})
                else:
                    body["paths"].append({"group": i, "tdn": x["tdn"]})
        return body

    def template(self, client, site, product_id):
        return next(p for p in client.get(f"/aci/api/templates?site={site}").get_json()["products"] if p["id"] == product_id)

    def forget_template_changes(self):
        """Back to the shipped templates after a test, whatever it left behind."""
        def clean():
            with sqlite3.connect(audit.DB_PATH) as db:
                db.execute("DELETE FROM aci_templates")
                db.execute("DELETE FROM aci_products")
        self.addCleanup(clean)

    def test_template_editing_needs_the_permission(self):
        self.forget_template_changes()
        operator, admin = self.client(), self.client("admin")
        self.assertFalse(operator.get("/aci/api/session").get_json()["can_edit_templates"])
        self.assertTrue(admin.get("/aci/api/session").get_json()["can_edit_templates"])
        flex = self.template(operator, "DCA", "flex")
        self.assertFalse(operator.get("/aci/api/templates?site=DCA").get_json()["can_edit"])
        body = self.editor_body(flex)
        body["groups"][0]["default"] = not body["groups"][0]["default"]
        for url, data in (("/aci/api/templates/flex/preview?site=DCA", body), ("/aci/api/templates/flex/save?site=DCA", body),
                          ("/aci/api/products", {"label": "X"}), ("/aci/api/products/flex/delete", {"confirm": True})):
            r = operator.post(url, json=data)
            self.assertEqual(r.status_code, 403, url)
            self.assertIn("ACI: edit product templates", r.get_json()["error"])
            self.assertEqual(self.app.test_client().post(url, json=data).status_code, 401, url)
        # The permission is ticked on the Users page, under the ACI module; "all modules" never includes it.
        users = admin.get("/api/admin/users").get_json()
        self.assertIn({"id": "aci-templates", "title": "ACI: edit product templates", "permission_of": "aci", "explicit": True}, users["modules"])
        r = admin.post("/api/admin/users", json={"username": "operator3", "role": "operator", "modules": ["aci", "aci-templates"]})
        self.assertEqual(r.status_code, 200, r.get_json())
        self.addCleanup(admin.delete, "/api/admin/users/operator3")
        editor = self.client("operator3")
        self.assertTrue(editor.get("/aci/api/session").get_json()["can_edit_templates"])
        result = editor.post("/aci/api/templates/flex/save?site=DCA", json=body).get_json()
        self.assertTrue(result["ok"], result)
        self.assertEqual(self.template(operator, "DCA", "flex")["updated_by"], "operator3")

    def test_template_save_adds_and_removes_paths_per_site(self):
        self.forget_template_changes()
        admin = self.client("admin")
        dcb_before = self.template(admin, "DCB", "flex")
        flex = self.template(admin, "DCA", "flex")
        self.assertEqual((flex["version"], flex["edited"], flex["custom"], flex["default_paths"]), (0, False, False, 13))
        body = self.editor_body(flex)
        removed = body["paths"].pop(0)["tdn"]
        body["physical"].append({"group": 0, "pod": "1", "node": "101", "interface": "eth1/13"})   # exists in the DCA fabric
        body["paths"].append({"group": 2, "tdn": "topology/pod-1/protpaths-103-104/pathep-[Openstack-Traffic-A-PG]"})
        posts, events = len(self.dca.posts), len(self.dca.requests)
        plan = admin.post("/aci/api/templates/flex/preview?site=DCA", json=body).get_json()
        self.assertTrue(plan["ok"], plan)
        self.assertEqual((plan["summary"]["new"], plan["summary"]["delete"], plan["summary"]["same"], plan["blocked"]), (2, 1, 12, False))
        self.assertTrue(all(i["exists"] for i in plan["items"]))
        self.assertEqual(self.template(admin, "DCA", "flex")["version"], 0, "a preview saves nothing")
        result = admin.post("/aci/api/templates/flex/save?site=DCA", json=body).get_json()
        self.assertTrue(result["ok"], result)
        self.assertEqual(result["summary_text"], "2 added, 1 removed (13 → 14 paths)")
        # Nothing was sent to APIC, only read from it.
        self.assertEqual(len(self.dca.posts), posts)
        self.assertTrue(all(method == "GET" or "aaa" in path for method, path in self.dca.requests[events:]), self.dca.requests[events:])
        after = self.template(admin, "DCA", "flex")
        tdns = [x["tdn"] for g in after["groups"] for x in g["paths"]]
        # Within a group the physical / port-channel rows come first, then the vPC / path TDN rows.
        self.assertEqual(after["groups"][0]["paths"][0]["tdn"], "topology/pod-1/paths-101/pathep-[eth1/13]")
        self.assertNotIn(removed, tdns)
        self.assertEqual((len(tdns), after["version"], after["edited"], after["updated_by"]), (14, 1, True, "admin"))
        self.assertEqual(self.template(admin, "DCB", "flex"), dcb_before, "another site keeps its template")
        # Static Port offers the edited template to every user.
        operator_view = self.template(self.client(), "DCA", "flex")
        self.assertEqual([x["tdn"] for g in operator_view["groups"] for x in g["paths"]], tdns)
        # History: who changed what, before and after; the activity log names it too.
        record = admin.get(f"/aci/api/audit/{result['audit_id']}").get_json()["item"]
        self.assertEqual((record["username"], record["site"], record["apic"], record["action"], record["product"], record["status"], record["target"],
                          record["apic_user"]), ("admin", "DCA", self.dca.host, "template", "Acme Flex", "success", "templates/DCA/flex", None))
        self.assertEqual(sorted((i["status"], i["tdn"]) for i in record["details"]["paths"]),
                         sorted([("new", "topology/pod-1/paths-101/pathep-[eth1/13]"), ("delete", removed),
                                 ("new", "topology/pod-1/protpaths-103-104/pathep-[Openstack-Traffic-A-PG]")]))
        self.assertEqual(sum(len(g["paths"]) for g in record["details"]["before"]), 13)
        self.assertEqual(admin.get("/aci/api/audit?action=template&site=DCA").get_json()["items"][0]["id"], result["audit_id"])
        event = admin.get("/api/admin/events?actor=admin&kinds=static_port&limit=1").get_json()["events"][0]
        self.assertEqual((event["action"], event["target"]), ("static_port.template_saved", "templates/DCA"))
        self.assertEqual(event["detail"], f"change #{result['audit_id']}, site DCA, product Acme Flex: 2 added, 1 removed (13 → 14 paths)")
        # A template change cannot be undone from History (it never touched APIC).
        r = admin.post(f"/aci/api/audit/{result['audit_id']}/undo/preview", json={})
        self.assertEqual(r.status_code, 400)
        self.assertIn("is a product template change", r.get_json()["error"])
        # Saving the same again changes nothing; an old version is refused.
        same = admin.post("/aci/api/templates/flex/save?site=DCA", json={**self.editor_body(after)})
        self.assertIn("Nothing to save", same.get_json()["error"])
        stale = admin.post("/aci/api/templates/flex/save?site=DCA", json=body)
        self.assertEqual(stale.status_code, 400)
        self.assertIn("was changed by admin", stale.get_json()["error"])
        # Reset to default: the shipped template again, recorded as a change of its own.
        plan = admin.post("/aci/api/templates/flex/preview?site=DCA", json={"version": 1, "reset": True}).get_json()
        self.assertEqual((plan["summary"]["new"], plan["summary"]["delete"], plan["reset"]), (1, 2, True))
        reset = admin.post("/aci/api/templates/flex/save?site=DCA", json={"version": 1, "reset": True}).get_json()
        self.assertTrue(reset["ok"], reset)
        default = self.template(admin, "DCA", "flex")
        self.assertEqual((default["edited"], default["version"], sum(len(g["paths"]) for g in default["groups"])), (False, 2, 13))
        self.assertTrue(admin.get(f"/aci/api/audit/{reset['audit_id']}").get_json()["item"]["summary"].startswith("reset to the default template: "))
        again = admin.post("/aci/api/templates/flex/save?site=DCA", json={"version": 2, "reset": True})
        self.assertIn("already the default", again.get_json()["error"])

    def test_template_paths_are_validated_and_looked_up_on_apic(self):
        self.forget_template_changes()
        admin = self.client("admin")
        prime = self.template(admin, "DCA", "prime")
        posts = len(self.dca.posts)
        body = self.editor_body(prime)
        body["physical"].append({"group": 0, "pod": "1", "node": "111", "interface": "eth1/99"})   # not in the fabric
        plan = admin.post("/aci/api/templates/prime/preview?site=DCA", json=body).get_json()
        self.assertTrue(plan["blocked"])
        self.assertEqual(plan["missing"], ["topology/pod-1/paths-111/pathep-[eth1/99]"])
        r = admin.post("/aci/api/templates/prime/save?site=DCA", json=body)
        self.assertEqual(r.status_code, 400)
        self.assertIn("do not exist in the DCA fabric, so the template was not saved: topology/pod-1/paths-111/pathep-[eth1/99]", r.get_json()["error"])
        self.assertFalse(self.template(admin, "DCA", "prime")["edited"])
        # A path already in the template that is gone from the fabric does not block other changes; it is only flagged.
        tdn = "topology/pod-1/paths-111/pathep-[eth1/1]"
        removed = self.dca.mit.pop(tdn)
        try:
            body = self.editor_body(self.template(admin, "DCA", "prime"))
            body["groups"][-1]["default"] = True
            plan = admin.post("/aci/api/templates/prime/preview?site=DCA", json=body).get_json()
            self.assertFalse(plan["blocked"], plan)
            self.assertFalse(next(i for i in plan["items"] if i["tdn"] == tdn)["exists"])
            self.assertEqual(plan["groups"], [{"name": "Uplink Metro Net & Router DCI (leaf 101-102)", "status": "change", "default": True}])
        finally:
            self.dca.mit[tdn] = removed
        # Invalid rows, an unknown product, a site without a service account.
        bad = {"version": 0, "groups": [{"name": "G"}], "physical": [{"group": 0, "node": "x", "interface": "eth1/1"}], "paths": [{"group": 0, "tdn": "uni/x"}]}
        r = admin.post("/aci/api/templates/prime/preview?site=DCA", json=bad)
        self.assertEqual(r.status_code, 400)
        self.assertIn("Physical/PC row 1: Node must be a number", r.get_json()["error"])
        self.assertIn("vPC/TDN row 1: Invalid path TDN", r.get_json()["error"])
        for url, message in (("/aci/api/templates/nope/preview?site=DCA", "Unknown product"), ("/aci/api/templates/prime/preview?site=DCC", "service account"),
                             ("/aci/api/templates/prime/preview?site=EVIL", "Unknown site")):
            r = admin.post(url, json=self.editor_body(prime))
            self.assertEqual(r.status_code, 400, url)
            self.assertIn(message, r.get_json()["error"])
        self.assertIn("version is missing", admin.post("/aci/api/templates/prime/preview?site=DCA", json={"groups": []}).get_json()["error"])
        self.assertEqual(len(self.dca.posts), posts)
        # The editor's live check: which paths exist in the fabric (read-only).
        r = admin.post("/aci/api/paths/check?site=DCA", json={"tdns": [tdn, "topology/pod-1/paths-111/pathep-[eth1/99]"]}).get_json()
        self.assertEqual(r["exists"], {tdn: True, "topology/pod-1/paths-111/pathep-[eth1/99]": False})
        self.assertEqual(admin.post("/aci/api/paths/check?site=DCA", json={"tdns": "x"}).status_code, 400)
        self.assertEqual(self.client().post("/aci/api/paths/check?site=DCA", json={"tdns": []}).status_code, 200)   # every ACI user
        self.assertEqual(self.app.test_client().post("/aci/api/paths/check?site=DCA", json={"tdns": []}).status_code, 401)

    def test_products_added_by_users(self):
        self.forget_template_changes()
        admin, operator = self.client("admin"), self.client()
        r = admin.post("/aci/api/products", json={"label": " Acme  GPU ", "platform": "Bare metal", "mode": "access", "immediacy": "lazy"}).get_json()
        self.assertTrue(r["ok"], r)
        self.assertEqual({k: r["product"][k] for k in ("id", "label", "platform", "mode", "immediacy", "custom")},
                         {"id": "acme-gpu", "label": "Acme GPU", "platform": "Bare metal", "mode": "untagged", "immediacy": "lazy", "custom": True})
        added = r["audit_id"]
        for site in ("DCB", "DCA"):
            gpu = self.template(operator, site, "acme-gpu")
            self.assertEqual((gpu["groups"], gpu["version"], gpu["edited"]), ([], 0, False))
        self.assertEqual([p["id"] for p in operator.get("/aci/api/templates?site=DCA").get_json()["products"]], ["flex", "prime", "rhoso", "acme-gpu"])
        for body, message in (({"label": "acme gpu"}, "already exists"), ({"label": "Acme Flex"}, "already exists"), ({"label": ""}, "Product name is required"),
                              ({"label": "X", "mode": "bogus"}, "Unknown mode"), ({"label": "Y", "immediacy": "later"}, "Immediacy is invalid")):
            r = admin.post("/aci/api/products", json=body)
            self.assertEqual(r.status_code, 400, body)
            self.assertIn(message, r.get_json()["error"])
        # Its template on DCA; DCB stays empty. It has no default to reset to.
        body = {"version": 0, "groups": [{"name": "GPU servers", "default": True}],
                "physical": [{"group": 0, "node": "106", "interface": "eth1/17"}], "paths": []}
        saved = admin.post("/aci/api/templates/acme-gpu/save?site=DCA", json=body).get_json()
        self.assertTrue(saved["ok"], saved)
        self.assertEqual(self.template(operator, "DCA", "acme-gpu")["groups"][0]["paths"], [{"tdn": "topology/pod-1/paths-106/pathep-[eth1/17]", "exists": True}])
        self.assertEqual(self.template(operator, "DCB", "acme-gpu")["groups"], [])
        r = admin.post("/aci/api/templates/acme-gpu/preview?site=DCA", json={"version": 1, "reset": True})
        self.assertIn("has no default template", r.get_json()["error"])
        # Static Port records the product's label with the change.
        result = self.audit_apply(operator, "Openstack-Traffic-P-PG", "1843", product="acme-gpu").get_json()
        self.assertEqual(operator.get(f"/aci/api/audit/{result['audit_id']}").get_json()["item"]["product"], "Acme GPU")
        # Settings, then delete with its templates; built-in products stay.
        r = admin.post("/aci/api/products/acme-gpu", json={"label": "Acme GPU Cloud", "platform": "Bare metal", "mode": "untagged", "immediacy": "lazy"}).get_json()
        self.assertTrue(r["ok"], r)
        self.assertEqual(self.template(operator, "DCA", "acme-gpu")["label"], "Acme GPU Cloud")
        self.assertIn("unchanged", admin.post("/aci/api/products/acme-gpu", json={"label": "Acme GPU Cloud", "platform": "Bare metal", "mode": "untagged",
                                                                                 "immediacy": "lazy"}).get_json()["error"])
        for url, body, message in (("/aci/api/products/flex", {"label": "Flex 2"}, "Built-in products"), ("/aci/api/products/flex/delete", {"confirm": True}, "Built-in products"),
                                   ("/aci/api/products/acme-gpu/delete", {}, "Confirm the deletion"), ("/aci/api/products/nope/delete", {"confirm": True}, "Unknown product")):
            r = admin.post(url, json=body)
            self.assertEqual(r.status_code, 400, url)
            self.assertIn(message, r.get_json()["error"])
        deleted = admin.post("/aci/api/products/acme-gpu/delete", json={"confirm": True}).get_json()
        self.assertTrue(deleted["ok"], deleted)
        self.assertNotIn("acme-gpu", [p["id"] for p in operator.get("/aci/api/templates?site=DCA").get_json()["products"]])
        record = admin.get(f"/aci/api/audit/{deleted['audit_id']}").get_json()["item"]
        self.assertEqual((record["site"], record["action"], record["product"], record["summary"]),
                         ("", "template", "Acme GPU Cloud", "product deleted: Acme GPU Cloud (paths per site: DCA 1)"))
        self.assertEqual(record["details"]["templates"]["DCA"][0]["paths"], ["topology/pod-1/paths-106/pathep-[eth1/17]"])
        self.assertEqual(admin.get(f"/aci/api/audit/{added}").get_json()["item"]["summary"], "product added: Acme GPU (Bare metal)")
        self.assertNotIn("", admin.get("/aci/api/audit").get_json()["sites"])   # a product change is for every site
        actions = [e["action"] for e in admin.get("/api/admin/events?actor=admin&kinds=static_port&limit=4").get_json()["events"]]
        self.assertEqual(actions, ["static_port.product_deleted", "static_port.product_changed", "static_port.template_saved", "static_port.product_added"])

    # ---- IP usage (Client Endpoints) ----
    def ip_rows(self, client, site="DCB", **filters):
        ip_usage._cache.clear()   # the endpoints are read again from the mock
        r = client.get("/aci/api/ip-usage", query_string={"site": site, **filters}).get_json()
        self.assertTrue(r["ok"], r)
        return {(x["epg"], x["subnet"]): x for x in r["items"]}

    def test_ip_usage_learned_and_free_addresses(self):
        client = self.client()
        posts = len(self.dcb.posts)
        rows = self.ip_rows(client, tenant="CLOUD-03")
        main = rows[("11.8.158.1", "11.8.158.0/24")]
        self.assertEqual((main["gateways"], main["scope"], main["size"], main["learned"], main["used_elsewhere"], main["free"]),
                         (["11.8.158.1"], "public", 254, 15, 1, 237))
        self.assertEqual(main["suggest"], [f"11.8.158.{h}" for h in (2, 3, 5, 7, 8, 9, 14, 15, 17, 18)])
        self.assertEqual(main["free_ranges"][:3], ["11.8.158.2-11.8.158.3", "11.8.158.5", "11.8.158.7-11.8.158.9"])
        e = next(x for x in main["endpoints"] if x["ip"] == "11.8.158.23")
        self.assertEqual((e["mac"], e["encap"], e["interface"], e["source"]), ("00:50:56:8B:8F:E7", "vlan-1604", "Pod-1/Node-118-119/VPC-25G-CLOUD-A-PG", "learned"))
        # The second EPG of the same bridge domain: the IPs of the first one are not free there either.
        dmz = rows[("11.8.158.1-DMZ", "11.8.158.0/24")]
        self.assertEqual((dmz["learned"], dmz["used_elsewhere"], dmz["free"]), (1, 15, 237))
        full = rows[("11.13.128.1", "11.13.128.0/28")]
        self.assertEqual((full["size"], full["learned"], full["free"], full["suggest"], full["free_ranges"]), (14, 12, 1, ["11.13.128.9"], ["11.13.128.9"]))
        self.assertEqual(rows[("11.15.64.1", "11.15.64.0/26")]["learned"], 0)   # the page hides it unless asked
        self.assertNotIn("PAN-TO-ASR9K-L2", {epg for epg, _ in rows}, "an EPG with only MAC addresses has no row")
        self.assertEqual(len(self.dcb.posts), posts, "read-only on APIC")
        # Every learned IP is kept; an IP that is not learned any more stays out of the suggestions for a while.
        with sqlite3.connect(audit.DB_PATH) as db:
            self.assertEqual(db.execute("SELECT count(*) FROM aci_ip_seen WHERE site = 'DCB' AND ip LIKE '11.8.158.%'").fetchone()[0], 16)
            db.execute("INSERT INTO aci_ip_seen VALUES ('DCB', '11.8.158.2', 'uni/tn-CLOUD-03/ap-Ip-Public-Shared/epg-old', 'AA:BB:CC:00:00:01', ?, ?)",
                       ("2026-10-01T10:00:00+07:00", (datetime.datetime.now().astimezone() - datetime.timedelta(days=3)).isoformat(timespec="seconds")))
        dn = "uni/tn-CLOUD-03/ap-Ip-Public-Shared/epg-11.8.158.1/cep-00:50:56:8B:FC:AB"
        removed = {k: self.dcb.mit.pop(k) for k in [k for k in self.dcb.mit if k == dn or k.startswith(dn + "/")]}
        try:
            main = self.ip_rows(client, tenant="CLOUD-03", app="Ip-Public-Shared", epg="11.8.158.1")[("11.8.158.1", "11.8.158.0/24")]
        finally:
            self.dcb.mit.update(removed)
            with sqlite3.connect(audit.DB_PATH) as db:
                db.execute("DELETE FROM aci_ip_seen WHERE ip = '11.8.158.2'")
        self.assertEqual((main["learned"], main["free"], main["seen_recently"]), (14, 238, ["11.8.158.2", "11.8.158.16"]))
        self.assertEqual(main["suggest"][:3], ["11.8.158.3", "11.8.158.5", "11.8.158.7"])
        # Per site, and the input is checked like the other APIs.
        self.assertIn(("11.8.160.1", "11.8.160.0/24"), self.ip_rows(client, site="DCA", tenant="CLOUD-03"))
        self.assertEqual(client.get("/aci/api/ip-usage?site=EVIL").status_code, 400)
        self.assertEqual(client.get("/aci/api/ip-usage?site=DCB&tenant=a/b").status_code, 400)
        self.assertEqual(self.app.test_client().get("/aci/api/ip-usage?site=DCB").status_code, 401)

    def test_ip_usage_export(self):
        client = self.client()
        ip_usage._cache.clear()
        r = client.get("/aci/export/ip-usage.csv?site=DCB&tenant=CLOUD-03&scope=public&learned=1")
        self.assertEqual(r.status_code, 200)
        self.assertIn("ACI_IP_Usage_DCB_", r.headers["Content-Disposition"])
        rows = list(csv.DictReader(io.StringIO(r.get_data(as_text=True))))
        mine = [x for x in rows if x["EPG"] == "11.8.158.1"]
        self.assertEqual(sorted({x["Status"] for x in mine}), ["free", "gateway", "learned"])
        self.assertEqual(sum(x["Status"] == "learned" for x in mine), 15)
        self.assertNotIn("11.15.64.1", {x["EPG"] for x in rows}, "learned=1: subnets without learned IPs are left out")
        self.assertEqual(self.app.test_client().get("/aci/export/ip-usage.csv?site=DCB").status_code, 401)

    # ---- provisioning ----
    def test_provisioning_new_objects_in_one_transaction(self):
        client = self.client()
        req = {"tenant": {"mode": "new", "name": "NEW-T", "descr": "new tenant"},
               "vrf": {"mode": "new", "name": "NEW-VRF"},
               "bd": {"mode": "new", "name": "NEW-BD", "arp_flood": True, "unknown_unicast": "flood",
                      "subnets": [{"ip": "10.249.10.1/24", "scope": "public"}], "l3outs": ["L3OUT-X"]},
               "app": {"mode": "new", "name": "NEW-AP"}, "epg": {"name": "NEW-EPG"},
               "domains": ["uni/vmmp-OpenStack/dom-DCB-OSP", "uni/phys-DCB-PHYS"], "deployment": "immediate", "resolution": "pre-provision"}
        preview = client.post("/aci/api/provision/preview?site=DCB", json=req).get_json()
        self.assertTrue(preview["ok"], preview)
        self.assertEqual(preview["steps"][0], {"object": "Tenant", "name": "NEW-T", "kind": "new", "action": "create"})
        posts = len(self.dcb.posts)
        result = client.post("/aci/api/provision/apply?site=DCB", json=req).get_json()
        self.assertTrue(result["ok"], result)
        self.assertEqual(len(self.dcb.posts), posts + 1)
        mit = self.dcb.mit
        for dn in ("uni/tn-NEW-T", "uni/tn-NEW-T/ctx-NEW-VRF", "uni/tn-NEW-T/BD-NEW-BD/subnet-[10.249.10.1/24]",
                   "uni/tn-NEW-T/BD-NEW-BD/rsctx", "uni/tn-NEW-T/BD-NEW-BD/rsBDToOut-L3OUT-X", "uni/tn-NEW-T/ap-NEW-AP/epg-NEW-EPG/rsbd",
                   "uni/tn-NEW-T/ap-NEW-AP/epg-NEW-EPG/rsdomAtt-[uni/vmmp-OpenStack/dom-DCB-OSP]"):
            self.assertIn(dn, mit)
        self.assertEqual(mit["uni/tn-NEW-T"][1]["descr"], "new tenant")
        self.assertEqual(mit["uni/tn-NEW-T/BD-NEW-BD"][1]["arpFlood"], "yes")
        self.assertEqual(mit["uni/tn-NEW-T/ap-NEW-AP/epg-NEW-EPG/rsdomAtt-[uni/phys-DCB-PHYS]"][1]["resImedcy"], "pre-provision")
        self.assertNotIn("uni/tn-NEW-T", self.dca.mit)
        record = client.get(f"/aci/api/audit/{result['audit_id']}").get_json()["item"]
        self.assertEqual((record["action"], record["site"], record["target"], record["status"], record["username"], record["apic_user"]),
                         ("provisioning", "DCB", "uni/tn-NEW-T/ap-NEW-AP/epg-NEW-EPG", "success", "operator2", SERVICE_USER))
        self.assertEqual(record["details"]["payload"], result["payload"])

    def test_provisioning_leaves_existing_objects_alone(self):
        client = self.client()
        req = {"tenant": {"mode": "existing", "name": "CLOUD-1", "descr": "do not use"}, "vrf": {"mode": "none"},
               "bd": {"mode": "existing", "name": "BD-2"}, "app": {"mode": "existing", "name": "MGMT-InBand"},
               "epg": {"name": "MGMT-VYOS"}, "domains": []}
        preview = client.post("/aci/api/provision/preview?site=DCA", json=req).get_json()
        self.assertEqual(preview["payload"]["fvTenant"]["attributes"], {"dn": "uni/tn-CLOUD-1"})
        self.assertTrue(any("moved to BD 'BD-2'" in w for w in preview["warnings"]), preview["warnings"])
        self.assertTrue(any("EPG 'MGMT-VYOS' already exists" in w for w in preview["warnings"]), preview["warnings"])

    def test_provisioning_errors(self):
        client = self.client()
        base = {"tenant": {"mode": "existing", "name": "DOES-NOT-EXIST"}, "app": {"mode": "new", "name": "AP"}, "epg": {"name": "EPG"}}
        r = client.post("/aci/api/provision/preview?site=DCB", json=base)
        self.assertEqual(r.status_code, 400)
        self.assertIn("does not exist on this site", r.get_json()["error"])
        bad_subnet = {**base, "tenant": {"mode": "new", "name": "T2"}, "bd": {"mode": "new", "name": "BD", "subnets": [{"ip": "10.253.0.0/24"}]}}
        r = client.post("/aci/api/provision/preview?site=DCB", json=bad_subnet)
        self.assertEqual(r.status_code, 400)
        self.assertIn("network address", r.get_json()["error"])

    # ---- audit ----
    def audit_apply(self, client, pg, vlan, **extra):
        req = {"tenant": "CLOUD-1", "app": "PUBLIC-CLOUD-OPENSTACK", "epg": "MGMT-VYOS", **extra,
               "paths": [{"tdn": f"topology/pod-1/protpaths-109-110/pathep-[{pg}]", "vlan": vlan}]}
        return client.post("/aci/api/static-port/apply?site=DCA", json=req)

    def test_audit_records_successful_apply(self):
        client = self.client()
        result = self.audit_apply(client, "Openstack-Traffic-N-PG", "1841", product="flex").get_json()
        self.assertTrue(result["ok"], result)
        record = client.get(f"/aci/api/audit/{result['audit_id']}").get_json()["item"]
        # The AD account of the operator, and the APIC service account the change was sent with.
        self.assertEqual((record["username"], record["apic_user"], record["site"], record["apic"], record["action"], record["product"],
                          record["status"]), ("operator2", SERVICE_USER, "DCA", self.dca.host, "static-port", "Acme Flex", "success"))
        self.assertEqual(record["target"], "uni/tn-CLOUD-1/ap-PUBLIC-CLOUD-OPENSTACK/epg-MGMT-VYOS")
        self.assertEqual(record["summary"], "1 new, 0 changed")
        self.assertEqual(record["details"]["items"][0]["encap"], "vlan-1841")
        self.assertTrue(record["finished"])
        listed = client.get("/aci/api/audit?user=operator2&limit=1").get_json()["items"][0]
        self.assertEqual((listed["id"], listed["apic_user"]), (result["audit_id"], SERVICE_USER))
        # The activity log (Users page) names the change, with a link to it in History.
        event = self.client("admin").get("/api/admin/events?actor=operator2&kinds=static_port&limit=1").get_json()["events"][0]
        self.assertEqual((event["action"], event["target"], event["client_ip"]),
                         ("static_port.applied", "uni/tn-CLOUD-1/ap-PUBLIC-CLOUD-OPENSTACK/epg-MGMT-VYOS", "127.0.0.1"))
        self.assertEqual(event["detail"], f"change #{result['audit_id']}, site DCA, product Acme Flex: 1 new, 0 changed (VLAN 1841)")

    def test_audit_records_failed_apply(self):
        client = self.client()
        self.dca.fail_posts = True
        try:
            response = self.audit_apply(client, "Openstack-Traffic-O-PG", "1842")
        finally:
            self.dca.fail_posts = False
        self.assertEqual(response.status_code, 502)
        items = client.get("/aci/api/audit?status=failed&site=DCA").get_json()["items"]
        self.assertTrue(items and "Simulated APIC failure" in items[0]["error"], items)
        event = self.client("admin").get("/api/admin/events?actor=operator2&kinds=static_port&limit=1").get_json()["events"][0]
        self.assertEqual(event["action"], "static_port.failed")
        self.assertIn(f"change #{items[0]['id']}, site DCA: 1 new, 0 changed (VLAN 1842); ", event["detail"])
        self.assertIn("Simulated APIC failure", event["detail"])

    def test_client_address_of_a_browser_on_the_bastion(self):
        """Behind nginx, the bastion forwarder's Remote Desktop details are logged with the address; otherwise ignored."""
        forwarded = {"X-Portal-Client": "rdp=10.248.30.40;user=INTRANET%5Cjdoe"}
        bastion = {"REMOTE_ADDR": "172.26.0.100"}
        client = self.client()
        self.assertEqual(client.get("/aci/api/session", headers=forwarded, environ_base=bastion).get_json()["client"], "172.26.0.100")
        self.webapp.PORTAL.trusted_proxy = True
        try:
            self.assertEqual(client.get("/aci/api/session", headers=forwarded, environ_base=bastion).get_json()["client"],
                             "172.26.0.100 (RDP 10.248.30.40, INTRANET\\jdoe)")
            req = {"tenant": "CLOUD-1", "app": "PUBLIC-CLOUD-OPENSTACK", "epg": "MGMT-VYOS",
                   "paths": [{"tdn": "topology/pod-1/protpaths-109-110/pathep-[Openstack-Traffic-Q-PG]", "vlan": "1847"}]}
            result = client.post("/aci/api/static-port/apply?site=DCA", json=req, headers=forwarded, environ_base=bastion).get_json()
        finally:
            self.webapp.PORTAL.trusted_proxy = False
        self.assertTrue(result["ok"], result)
        record = client.get(f"/aci/api/audit/{result['audit_id']}").get_json()["item"]
        self.assertEqual(record["client_ip"], "172.26.0.100 (RDP 10.248.30.40, INTRANET\\jdoe)")

    def test_audit_marks_dropped_connection_as_unknown(self):
        client = self.client()
        self.dca.drop_posts = True
        try:
            response = self.audit_apply(client, "Openstack-Traffic-S-PG", "1846")
        finally:
            self.dca.drop_posts = False
        self.assertEqual(response.status_code, 502)
        items = client.get("/aci/api/audit?status=unknown&site=DCA").get_json()["items"]
        self.assertTrue(items and "APIC call failed" in items[0]["error"], items)

    def test_apply_blocked_when_audit_cannot_be_written(self):
        client = self.client()
        posts, original = len(self.dca.posts), audit.DB_PATH
        audit.DB_PATH = self.tmp  # a directory, cannot be opened as a database
        try:
            response = self.audit_apply(client, "Openstack-Traffic-Q-PG", "1844")
        finally:
            audit.DB_PATH = original
        self.assertEqual(response.status_code, 502)
        self.assertIn("Cannot write the audit log", response.get_json()["error"])
        self.assertEqual(len(self.dca.posts), posts, "without an audit record nothing may change on APIC")

    def test_audit_requires_login_and_validates_filters(self):
        anonymous = self.app.test_client()
        self.assertEqual(anonymous.get("/aci/api/audit").status_code, 401)
        self.assertEqual(anonymous.get("/aci/api/audit/1").status_code, 401)
        self.assertEqual(anonymous.get("/aci/export/audit.csv").status_code, 401)
        client = self.client()
        self.assertEqual(client.get("/aci/api/audit?since=kemarin").status_code, 400)
        self.assertEqual(client.get("/aci/api/audit?action=remove").status_code, 400)
        self.assertEqual(client.get("/aci/api/audit?limit=abc").status_code, 400)
        self.audit_apply(client, "Openstack-Traffic-R-PG", "1845")
        data = client.get("/aci/api/audit?user=operator2&limit=1").get_json()
        self.assertEqual(len(data["items"]), 1)
        self.assertIn("operator2", data["users"])
        self.assertIn("DCA", data["sites"])

    def test_audit_csv_neutralises_spreadsheet_formulas(self):
        # AD usernames cannot look like formulas, but APIC usernames and records of older versions can.
        record_id = audit.begin("=HYPERLINK(1)", "DCA", self.dca.host, "static-port", "uni/tn-CLOUD-1/ap-A/epg-B", None,
                                "1 new, 0 changed", {"items": []}, "127.0.0.1", apic_user="@SUM(1)")
        audit.finish(record_id, "success")
        csv_text = self.client().get("/aci/export/audit.csv", query_string={"user": "=HYPERLINK(1)"}).data.decode()
        self.assertTrue(csv_text.startswith("ID,Time,"))
        self.assertIn("'=HYPERLINK(1)", csv_text)
        self.assertNotIn(",=HYPERLINK", csv_text)
        self.assertIn(",'@SUM(1)", csv_text)

    # ---- report ----
    def test_report_and_exports(self):
        client = self.client()
        rows = client.get("/aci/api/report?site=DCB&tenant=CLOUD-1").get_json()["items"]
        self.assertEqual(sorted(r["Binding Type"] for r in rows), ["PC", "Physical", "Physical", "vPC"])
        self.assertTrue(all(r["Site"] == "DCB" for r in rows))
        r = client.get("/aci/export/report.xlsx?site=DCB&tenant=CLOUD-1")
        self.assertEqual(r.status_code, 200)
        self.assertRegex(r.headers["Content-Disposition"], r"ACI_Static_Port_Report_DCB_\d{8}-\d{4}\.xlsx")
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(r.data))
        summary = {row[0]: row[1] for row in workbook["Summary"].iter_rows(min_row=2, values_only=True)}
        self.assertTrue(summary["Site"].startswith("DCB"))
        self.assertEqual((summary["Filter"], summary["Total Bindings"], summary["PC Bindings"], summary["Generated by"]),
                         ("tenant=CLOUD-1", 4, 1, "operator2"))
        self.assertEqual(workbook["Static Port Report"].max_row, 5)
        csv_text = client.get("/aci/export/report.csv?site=DCA&tenant=CLOUD-1&app=MGMT-InBand&epg=Mgmt-InBand-EPG").data.decode()
        self.assertTrue(csv_text.startswith("Site,Tenant,Application Profile"))
        self.assertEqual(client.get("/aci/export/report.xlsx?site=EVIL").status_code, 400)
        self.assertEqual(self.app.test_client().get("/aci/export/report.csv?site=DCB").status_code, 401)

    def test_report_paging(self):
        client = self.client()
        page_size, aci.PAGE_SIZE = aci.PAGE_SIZE, 3
        try:
            rows = client.get("/aci/api/report?site=DCB&tenant=CLOUD-1").get_json()["items"]
        finally:
            aci.PAGE_SIZE = page_size
        self.assertEqual(len(rows), 4)

    def test_report_shows_where_each_binding_comes_from(self):
        client = self.client()
        epg_dn = self.make_epg(self.dca, "DRIFT-EPG")
        path = "topology/pod-1/paths-101/pathep-[eth1/{}]".format
        req = {"tenant": "CLOUD-1", "app": "MGMT-InBand", "epg": "DRIFT-EPG",
               "physical": [{"node": "101", "interface": f"eth1/{n}", "vlan": "3101"} for n in (51, 52, 53)]}
        change = client.post("/aci/api/static-port/apply?site=DCA", json=req).get_json()["audit_id"]
        # Outside the portal: a new binding, another VLAN on a binding of the portal, and a binding removed.
        self.bind(self.dca, epg_dn, path(54), 3102)
        self.dca.mit[f"{epg_dn}/rspathAtt-[{path(52)}]"][1]["encap"] = "vlan-3103"
        del self.dca.mit[f"{epg_dn}/rspathAtt-[{path(53)}]"]
        # A path the portal deleted (with an undo) that is back again.
        req55 = dict(req, physical=[{"node": "101", "interface": "eth1/55", "vlan": "3104"}])
        second = client.post("/aci/api/static-port/apply?site=DCA", json=req55).get_json()["audit_id"]
        undo = client.post(f"/aci/api/audit/{second}/undo", json={"confirm": True}).get_json()["audit_id"]
        self.bind(self.dca, epg_dn, path(55), 3104)
        query = "site=DCA&tenant=CLOUD-1&app=MGMT-InBand&epg=DRIFT-EPG"
        rows = {r["Path TDN"]: r for r in client.get(f"/aci/api/report?{query}").get_json()["items"]}
        self.assertEqual({tdn: (r["source"], r["Source"]) for tdn, r in rows.items()},
                         {path(51): ("portal", "Portal"), path(52): ("changed", "Changed outside portal"),
                          path(53): ("removed", "Removed outside portal"), path(54): ("outside", "Outside portal"),
                          path(55): ("changed", "Changed outside portal")})
        self.assertEqual(rows[path(51)]["portal"], {"change": change, "status": "new", "encap": "vlan-3101", "mode": "regular",
                                                    "immediacy": "immediate"})
        self.assertEqual((rows[path(52)]["VLAN"], rows[path(52)]["portal"]["encap"]), (3103, "vlan-3101"))   # APIC now / portal
        self.assertEqual((rows[path(55)]["portal"]["change"], rows[path(55)]["portal"]["status"]), (undo, "delete"))
        self.assertIsNone(rows[path(54)]["portal"])
        # The removed binding: the portal's values, no APIC data.
        removed = rows[path(53)]
        self.assertEqual((removed["VLAN"], removed["Encapsulation"], removed["Mode"], removed["Binding Type"], removed["Node(s)"],
                          removed["Last Changed (APIC)"], removed["Site"]), (3101, "vlan-3101", "regular", "Physical", "101", "", "DCA"))
        self.assertTrue(rows[path(51)]["Last Changed (APIC)"])   # modTs of the binding on APIC
        self.assertEqual(rows[path(54)]["Last Changed (APIC)"], "")   # not in the mock for a binding made outside
        # Only the report's own filters count: the whole site lists these paths once, another app's report none of them.
        everything = client.get("/aci/api/report?site=DCA").get_json()["items"]
        self.assertEqual(sorted(r["Path TDN"] for r in everything if r["EPG"] == "DRIFT-EPG"), sorted(rows))
        tenant = client.get("/aci/api/report?site=DCA&tenant=CLOUD-1").get_json()["items"]
        self.assertEqual({r["Path TDN"]: r["source"] for r in tenant if r["EPG"] == "DRIFT-EPG"}, {t: r["source"] for t, r in rows.items()})
        other = client.get("/aci/api/report?site=DCA&tenant=CLOUD-1&app=PUBLIC-CLOUD-OPENSTACK").get_json()["items"]
        self.assertFalse([r for r in other if r["EPG"] == "DRIFT-EPG"])
        # The same EPG on another site has no portal changes: its bindings there are outside the portal.
        dcb_epg = self.make_epg(self.dcb, "DRIFT-EPG")
        self.bind(self.dcb, dcb_epg, path(51), 3101)
        dcb = client.get("/aci/api/report?site=DCB&tenant=CLOUD-1&app=MGMT-InBand&epg=DRIFT-EPG").get_json()["items"]
        self.assertEqual([(r["Path TDN"], r["source"]) for r in dcb], [(path(51), "outside")])
        # Exports: source, portal change and setting, and the time of the last change on APIC.
        header, *lines = csv.reader(io.StringIO(client.get(f"/aci/export/report.csv?{query}").data.decode()))
        self.assertEqual(header[-4:], ["Source", "Portal Change", "Portal Setting", "Last Changed (APIC)"])
        exported = {line[header.index("Path TDN")]: dict(zip(header, line)) for line in lines}
        self.assertEqual([exported[path(n)]["Source"] for n in (51, 52, 53, 54, 55)],
                         ["Portal", "Changed outside portal", "Removed outside portal", "Outside portal", "Changed outside portal"])
        self.assertEqual((exported[path(52)]["Portal Change"], exported[path(52)]["Portal Setting"]), (str(change), "vlan-3101, regular, immediate"))
        self.assertEqual((exported[path(55)]["Portal Change"], exported[path(55)]["Portal Setting"]), (str(undo), "deleted"))
        self.assertEqual((exported[path(54)]["Portal Change"], exported[path(54)]["Portal Setting"], exported[path(54)]["Last Changed (APIC)"]), ("", "", ""))
        self.assertRegex(exported[path(51)]["Last Changed (APIC)"], r"^\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2}$")
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(client.get(f"/aci/export/report.xlsx?{query}").data))
        summary = {row[0]: row[1] for row in workbook["Summary"].iter_rows(min_row=2, values_only=True)}
        self.assertEqual([summary[f"Source: {s}"] for s in ("Portal", "Outside portal", "Changed outside portal", "Removed outside portal")],
                         [1, 1, 2, 1])
        sheet = workbook["Static Port Report"]
        self.assertEqual([c.value for c in sheet[1]][-4:], ["Source", "Portal Change", "Portal Setting", "Last Changed (APIC)"])
        self.assertEqual(sheet.max_row, 6)

    # ---- undo of a static-port change ----
    def test_undo_deletes_new_paths_and_restores_changed_ones(self):
        client, admin = self.client(), self.client("admin")
        epg_dn = self.make_epg(self.dca, "UNDO-EPG")
        path = "topology/pod-1/paths-101/pathep-[eth1/{}]".format
        self.bind(self.dca, epg_dn, path(61), 3001)   # the binding before the change
        req = {"tenant": "CLOUD-1", "app": "MGMT-InBand", "epg": "UNDO-EPG", "product": "flex",
               "physical": [{"node": "101", "interface": "eth1/61", "vlan": "3002", "mode": "native"},
                            {"node": "101", "interface": "eth1/62", "vlan": "3002"}]}
        change = client.post("/aci/api/static-port/apply?site=DCA", json=req).get_json()["audit_id"]
        # The undo always runs on the site of the change (DCA), whatever site the browser has selected.
        posts, dcb_posts = len(self.dca.posts), len(self.dcb.posts)
        preview = client.post(f"/aci/api/audit/{change}/undo/preview?site=DCB", json={}).get_json()
        self.assertTrue(preview["ok"], preview)
        self.assertEqual((preview["site"], preview["undo_of"], preview["epg_dn"], preview["summary"], preview["blocked"], preview["conflicts"]),
                         ("DCA", change, epg_dn, {"delete": 1, "restore": 1, "skip": 0}, False, []))
        items = {i["tdn"]: i for i in preview["items"]}
        self.assertEqual((items[path(61)]["action"], items[path(61)]["encap"], items[path(61)]["mode"], items[path(61)]["before"]["encap"],
                          items[path(61)]["before"]["mode"]), ("restore", "vlan-3001", "regular", "vlan-3002", "native"))
        self.assertEqual((items[path(62)]["action"], items[path(62)]["before"]["encap"]), ("delete", "vlan-3002"))
        self.assertEqual(preview["payload"], {"fvAEPg": {"attributes": {"dn": epg_dn}, "children": [
            {"fvRsPathAtt": {"attributes": {"tDn": path(61), "encap": "vlan-3001", "mode": "regular", "instrImedcy": "immediate"}}},
            {"fvRsPathAtt": {"attributes": {"tDn": path(62), "status": "deleted"}}}]}})
        self.assertEqual(len(self.dca.posts), posts, "a preview must not change APIC")
        r = client.post(f"/aci/api/audit/{change}/undo", json={})
        self.assertEqual(r.status_code, 400)
        self.assertIn("Confirm the undo", r.get_json()["error"])
        result = client.post(f"/aci/api/audit/{change}/undo?site=DCB", json={"confirm": True}).get_json()
        self.assertTrue(result["ok"], result)
        self.assertEqual((result["applied"], result["site"], result["undo_of"]), (True, "DCA", change))
        self.assertEqual((len(self.dca.posts), len(self.dcb.posts)), (posts + 1, dcb_posts), "one POST, to DCA only")
        restored = self.dca.mit[f"{epg_dn}/rspathAtt-[{path(61)}]"][1]
        self.assertEqual((restored["encap"], restored["mode"], restored["instrImedcy"]), ("vlan-3001", "regular", "immediate"))
        self.assertNotIn(f"{epg_dn}/rspathAtt-[{path(62)}]", self.dca.mit)
        # The undo is a change of its own, and both records point to each other.
        undo_id = result["audit_id"]
        undo = client.get(f"/aci/api/audit/{undo_id}").get_json()["item"]
        self.assertEqual((undo["action"], undo["site"], undo["target"], undo["status"], undo["username"], undo["product"],
                          undo["undo_of"], undo["undone_by"], undo["summary"]),
                         ("static-port", "DCA", epg_dn, "success", "operator2", "Acme Flex", change, None,
                          f"undo of change #{change}: 1 removed, 1 restored"))
        self.assertEqual(undo["details"], {"undo_of": change, "items": [
            {"tdn": path(61), "encap": "vlan-3001", "mode": "regular", "immediacy": "immediate", "status": "change",
             "before": {"tDn": path(61), "encap": "vlan-3002", "mode": "native", "instrImedcy": "immediate"}},
            {"tdn": path(62), "status": "delete", "before": {"tDn": path(62), "encap": "vlan-3002", "mode": "regular", "instrImedcy": "immediate"}}]})
        original = client.get(f"/aci/api/audit/{change}").get_json()["item"]
        self.assertEqual((original["undone_by"], original["undo_of"]), (undo_id, None))
        listed = {x["id"]: x for x in client.get("/aci/api/audit", query_string={"q": "UNDO-EPG"}).get_json()["items"]}
        self.assertEqual({i: (x["undo_of"], x["undone_by"]) for i, x in listed.items()}, {change: (None, undo_id), undo_id: (change, None)})
        rows = {int(r["ID"]): r for r in csv.DictReader(io.StringIO(client.get("/aci/export/audit.csv", query_string={"q": "UNDO-EPG"}).data.decode()))}
        self.assertEqual({i: (r["Undo of"], r["Undone by"]) for i, r in rows.items()}, {change: ("", str(undo_id)), undo_id: (str(change), "")})
        # Activity log: its own action, naming both changes and the VLANs (the one put back and the one of the deleted path).
        event = admin.get("/api/admin/events?actor=operator2&kinds=static_port&limit=1").get_json()["events"][0]
        self.assertEqual((event["action"], event["target"]), ("static_port.undone", epg_dn))
        self.assertEqual(event["detail"], f"change #{undo_id}, site DCA, product Acme Flex: undo of change #{change}: 1 removed, 1 restored"
                                          " (VLAN 3001, 3002)")
        # The report reads the undo: the restored path belongs to the portal again, the deleted path is gone as it should be.
        report = client.get("/aci/api/report?site=DCA&tenant=CLOUD-1&app=MGMT-InBand&epg=UNDO-EPG").get_json()["items"]
        self.assertEqual([(r["Path TDN"], r["source"], r["portal"]["change"]) for r in report], [(path(61), "portal", undo_id)])
        # Neither can be undone (again).
        for record_id, body, message in ((change, {}, f"already undone by change #{undo_id}"), (change, {"confirm": True}, "already undone"),
                                         (undo_id, {}, "is itself an undo"), (undo_id, {"confirm": True}, "is itself an undo")):
            r = client.post(f"/aci/api/audit/{record_id}/undo{'' if body else '/preview'}", json=body)
            self.assertEqual(r.status_code, 400, (record_id, body))
            self.assertIn(message, r.get_json()["error"])
        self.assertEqual(len(self.dca.posts), posts + 1)

    def test_undo_skips_paths_changed_or_removed_since(self):
        client = self.client()
        epg_dn = self.make_epg(self.dca, "UNDO-SKIP")
        path = "topology/pod-1/paths-102/pathep-[eth1/{}]".format
        self.bind(self.dca, epg_dn, path(4), 3010)
        req = {"tenant": "CLOUD-1", "app": "MGMT-InBand", "epg": "UNDO-SKIP",
               "physical": [{"node": "102", "interface": f"eth1/{n}", "vlan": "3011"} for n in (1, 2, 3, 4)]}
        change = client.post("/aci/api/static-port/apply?site=DCA", json=req).get_json()["audit_id"]
        self.dca.mit[f"{epg_dn}/rspathAtt-[{path(1)}]"][1]["encap"] = "vlan-3012"   # changed since
        del self.dca.mit[f"{epg_dn}/rspathAtt-[{path(2)}]"]                          # a new path, removed since
        del self.dca.mit[f"{epg_dn}/rspathAtt-[{path(4)}]"]                          # a changed path, removed since
        preview = client.post(f"/aci/api/audit/{change}/undo/preview", json={}).get_json()
        self.assertEqual((preview["summary"], preview["blocked"]), ({"delete": 1, "restore": 0, "skip": 3}, False))
        items = {i["tdn"]: i for i in preview["items"]}
        self.assertEqual({t: i["action"] for t, i in items.items()}, {path(1): "skip", path(2): "skip", path(3): "delete", path(4): "skip"})
        self.assertEqual(items[path(1)]["reason"], "Changed since: APIC now has vlan-3012, regular, immediate; a later change is never overwritten.")
        self.assertEqual(items[path(1)]["before"]["encap"], "vlan-3012")   # the values on APIC now
        self.assertTrue(items[path(2)]["reason"].startswith("Already removed"))
        self.assertTrue(items[path(4)]["reason"].startswith("Removed since"))
        self.assertIsNone(items[path(4)]["before"])
        self.assertEqual(preview["payload"]["fvAEPg"]["children"], [{"fvRsPathAtt": {"attributes": {"tDn": path(3), "status": "deleted"}}}])
        result = client.post(f"/aci/api/audit/{change}/undo", json={"confirm": True}).get_json()
        self.assertTrue(result["ok"], result)
        self.assertEqual(self.dca.mit[f"{epg_dn}/rspathAtt-[{path(1)}]"][1]["encap"], "vlan-3012")   # the later change is kept
        self.assertNotIn(f"{epg_dn}/rspathAtt-[{path(3)}]", self.dca.mit)
        undo = client.get(f"/aci/api/audit/{result['audit_id']}").get_json()["item"]
        self.assertEqual(undo["summary"], f"undo of change #{change}: 1 removed, 0 restored, 3 skipped")
        self.assertEqual([i["tdn"] for i in undo["details"]["items"]], [path(3)])
        self.assertEqual([s["tdn"] for s in undo["details"]["skipped"]], [path(1), path(2), path(4)])
        # When every path was changed since, there is nothing to undo: refused, nothing sent, nothing recorded.
        again = client.post("/aci/api/static-port/apply?site=DCA", json=dict(req, physical=[{"node": "102", "interface": "eth1/5", "vlan": "3013"}]))
        second = again.get_json()["audit_id"]
        self.dca.mit[f"{epg_dn}/rspathAtt-[{path(5)}]"][1]["mode"] = "native"
        preview = client.post(f"/aci/api/audit/{second}/undo/preview", json={}).get_json()
        self.assertEqual((preview["summary"], preview["payload"]), ({"delete": 0, "restore": 0, "skip": 1}, None))
        posts, records = len(self.dca.posts), audit.query({}, 1, 0)[0]
        r = client.post(f"/aci/api/audit/{second}/undo", json={"confirm": True})
        self.assertEqual(r.status_code, 400)
        self.assertIn(f"Nothing to undo: every path of change #{second} was changed or removed since", r.get_json()["error"])
        self.assertEqual((len(self.dca.posts), audit.query({}, 1, 0)[0]), (posts, records))

    def test_undo_blocked_by_a_vlan_another_epg_uses_now(self):
        client = self.client()
        epg_dn = self.make_epg(self.dca, "UNDO-VLAN")
        tdn = "topology/pod-1/paths-103/pathep-[eth1/7]"
        self.bind(self.dca, epg_dn, tdn, 3021)
        req = {"tenant": "CLOUD-1", "app": "MGMT-InBand", "epg": "UNDO-VLAN", "physical": [{"node": "103", "interface": "eth1/7", "vlan": "3022"}]}
        change = client.post("/aci/api/static-port/apply?site=DCA", json=req).get_json()["audit_id"]
        # Meanwhile another EPG took VLAN 3021: putting it back would bind it twice, so the undo is blocked like an apply.
        other = self.bind(self.dca, self.make_epg(self.dca, "UNDO-OTHER"), "topology/pod-1/paths-103/pathep-[eth1/8]", 3021)
        preview = client.post(f"/aci/api/audit/{change}/undo/preview", json={}).get_json()
        self.assertEqual((preview["blocked"], preview["payload"], [c["vlan"] for c in preview["conflicts"]]), (True, None, [3021]))
        used_by = preview["conflicts"][0]["used_by"][0]
        self.assertEqual((used_by["name"], used_by["binding"], used_by["same_leaf"]), ("UNDO-OTHER", "static path", True))
        self.assertTrue(preview["items"][0]["conflict"])
        posts, records = len(self.dca.posts), audit.query({}, 1, 0)[0]
        r = client.post(f"/aci/api/audit/{change}/undo", json={"confirm": True})
        self.assertEqual(r.status_code, 400)
        self.assertIn("VLAN 3021 is already used by EPG CLOUD-1/MGMT-InBand/UNDO-OTHER", r.get_json()["error"])
        self.assertIn("The undo cannot put these VLANs back", r.get_json()["error"])
        self.assertEqual((len(self.dca.posts), audit.query({}, 1, 0)[0]), (posts, records), "nothing sent, nothing recorded")
        self.assertIsNone(client.get(f"/aci/api/audit/{change}").get_json()["item"]["undone_by"])
        # Once the VLAN is free again, the undo goes through.
        del self.dca.mit[other]
        self.assertTrue(client.post(f"/aci/api/audit/{change}/undo", json={"confirm": True}).get_json()["ok"])
        self.assertEqual(self.dca.mit[f"{epg_dn}/rspathAtt-[{tdn}]"][1]["encap"], "vlan-3021")

    def test_undo_is_for_administrators_and_the_user_who_made_the_change(self):
        admin = self.client("admin")
        r = admin.post("/api/admin/users", json={"username": "operator3", "role": "operator", "modules": ["aci"]})
        self.assertEqual(r.status_code, 200, r.get_json())
        self.addCleanup(admin.delete, "/api/admin/users/operator3")
        owner, other = self.client(), self.client("operator3")
        epg_dn = self.make_epg(self.dca, "UNDO-PERM")
        req = {"tenant": "CLOUD-1", "app": "MGMT-InBand", "epg": "UNDO-PERM", "physical": [{"node": "104", "interface": "eth1/9", "vlan": "3031"}]}
        change = owner.post("/aci/api/static-port/apply?site=DCA", json=req).get_json()["audit_id"]
        posts = len(self.dca.posts)
        for url, body in ((f"/aci/api/audit/{change}/undo/preview", {}), (f"/aci/api/audit/{change}/undo", {"confirm": True})):
            r = other.post(url, json=body)
            self.assertEqual(r.status_code, 403, url)
            self.assertIn(f"Only administrators and operator2, who made change #{change}, can undo it", r.get_json()["error"])
            self.assertEqual(self.app.test_client().post(url, json=body).status_code, 401)
        self.assertEqual(len(self.dca.posts), posts)
        # An administrator may undo anyone's change; the undo is recorded under the administrator's account.
        self.assertTrue(admin.post(f"/aci/api/audit/{change}/undo/preview", json={}).get_json()["ok"])
        result = admin.post(f"/aci/api/audit/{change}/undo", json={"confirm": True}).get_json()
        self.assertTrue(result["ok"], result)
        self.assertEqual(admin.get(f"/aci/api/audit/{result['audit_id']}").get_json()["item"]["username"], "admin")
        self.assertNotIn(f"{epg_dn}/rspathAtt-[topology/pod-1/paths-104/pathep-[eth1/9]]", self.dca.mit)

    def test_undo_is_refused_for_provisioning_and_unsuccessful_changes(self):
        client = self.client()
        target = "uni/tn-CLOUD-1/ap-MGMT-InBand/epg-MGMT-VYOS"
        item = {"tdn": "topology/pod-1/paths-101/pathep-[eth1/70]", "encap": "vlan-3041", "mode": "regular", "immediacy": "immediate",
                "status": "new", "before": None}
        records = {}
        for action, status, details in (("provisioning", "success", {"steps": [], "warnings": [], "payload": {}}),
                                        ("static-port", "failed", {"items": [item]}), ("static-port", "unknown", {"items": [item]})):
            records[action, status] = audit.begin("operator2", "DCA", self.dca.host, action, target, None, "1 created, 0 updated",
                                                  details, "127.0.0.1")
            audit.finish(records[action, status], status, "APIC HTTP 400" if status == "failed" else None)
        posts = len(self.dca.posts)
        for record_id, message in ((records["provisioning", "success"], "Provisioning changes cannot be undone here, because the portal never deletes objects"),
                                   (records["static-port", "failed"], "Only a successful change can be undone"),
                                   (records["static-port", "unknown"], "has status unknown"),
                                   (999999, "Audit record #999999 not found")):
            for suffix, body in (("/preview", {}), ("", {"confirm": True})):
                r = client.post(f"/aci/api/audit/{record_id}/undo{suffix}", json=body)
                self.assertEqual(r.status_code, 400, (record_id, suffix))
                self.assertIn(message, r.get_json()["error"])
        self.assertEqual(len(self.dca.posts), posts)

    # ---- APIC tokens of the service accounts ----
    def test_token_refresh_keepalive_and_relogin(self):
        client = self.client()
        for site in ("DCB", "DCA"):
            self.webapp.CONNECTIONS.drop(site)
            self.assertEqual(client.get(f"/aci/api/discover?kind=tenants&site={site}").status_code, 200)   # fresh tokens
        time.sleep(REFRESH_TIMEOUT / 2 + 0.3)
        refreshed = self.dcb.count("GET", "/api/aaaRefresh.json")
        self.assertEqual(client.get("/aci/api/discover?kind=tenants&site=DCB").status_code, 200)
        self.assertEqual(self.dcb.count("GET", "/api/aaaRefresh.json"), refreshed + 1)
        refreshed, logins = self.dca.count("GET", "/api/aaaRefresh.json"), self.dca.count("POST", "/api/aaaLogin.json")
        self.assertIsNotNone(client.post("/aci/api/keepalive").get_json()["user"])
        self.assertEqual(self.dca.count("GET", "/api/aaaRefresh.json"), refreshed + 1)
        time.sleep(REFRESH_TIMEOUT / 2 + 0.3)  # the first DCA token is dead by now; the refreshed one keeps DCA going
        self.assertEqual(client.get("/aci/api/discover?kind=tenants&site=DCA").status_code, 200)
        self.assertEqual(self.dca.count("POST", "/api/aaaLogin.json"), logins)
        # Idle for longer than the token lifetime: the next request logs in again by itself.
        time.sleep(REFRESH_TIMEOUT + 0.5)
        r = client.get("/aci/api/discover?kind=tenants&site=DCA")
        self.assertEqual(r.status_code, 200, r.get_json())
        self.assertEqual(self.dca.count("POST", "/api/aaaLogin.json"), logins + 1)
        # keepalive never logs in by itself, and without a session it only reports the status.
        self.webapp.CONNECTIONS.drop("DCB")
        logins = self.dcb.count("POST", "/api/aaaLogin.json")
        client.post("/aci/api/keepalive")
        self.assertIsNone(self.app.test_client().post("/aci/api/keepalive").get_json()["user"])
        self.assertEqual(self.dcb.count("POST", "/api/aaaLogin.json"), logins)


if __name__ == "__main__":
    unittest.main()
